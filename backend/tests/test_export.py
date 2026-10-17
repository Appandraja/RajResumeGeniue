"""PDF / DOCX export: headers, extractable text, ATS structure, ownership, edge cases."""

import io
import zipfile
from collections.abc import Callable
from typing import Any

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader
from sqlalchemy.orm import Session

from app.models import Template
from app.services import export_service
from tests.conftest import resume_payload

Headers = dict[str, str]
MakeResume = Callable[..., dict[str, Any]]

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _pdf_text(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _export(client: TestClient, headers: Headers, rid: int, fmt: str) -> Any:
    return client.get(f"/api/resumes/{rid}/export/{fmt}", headers=headers)


def _norm(text: str) -> str:
    return " ".join(text.split())


# --------------------------------------------------------------------- PDF
def test_pdf_export_headers(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(auth_headers)
    response = _export(client, auth_headers, created["id"], "pdf")
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")
    assert response.headers["content-type"] == "application/pdf"
    assert (
        response.headers["content-disposition"]
        == 'attachment; filename="Backend_Engineer_Resume.pdf"'
    )
    assert response.headers["cache-control"] == "no-store"


def test_pdf_text_is_extractable(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(auth_headers)
    text = _norm(_pdf_text(_export(client, auth_headers, created["id"], "pdf").content))
    for expected in (
        "Alice Example",
        "alice@example.com",
        "Berlin, Germany",
        "Senior Engineer",
        "Acme",
        "Jan 2020 - Jun 2023",
        "Built the billing API",
        "Cut p99 latency by 40%",
        "BSc Computer Science",
        "Python, FastAPI, PostgreSQL, Docker, AWS",
    ):
        assert expected in text, expected


def test_pdf_single_column_standard_heading_order(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(
        auth_headers,
        sections=[
            {"type": "certification", "title": "AWS Certified", "start_date": "2021-05-01"},
            {"type": "project", "title": "Open Source Tool", "bullets": ["Wrote it"]},
            {"type": "education", "title": "BSc", "start_date": "2012-09-01"},
            {"type": "experience", "title": "Engineer", "start_date": "2020-01-01"},
        ],
    )
    text = _pdf_text(_export(client, auth_headers, created["id"], "pdf").content)
    order = [text.index(h) for h in ("Summary", "Experience", "Education", "Skills")]
    assert order == sorted(order)
    assert text.index("Skills") < text.index("Projects") < text.index("Certifications")
    # Contact block comes first, before any heading.
    assert text.index("alice@example.com") < text.index("Summary")


def test_pdf_entries_reverse_chronological_and_present(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(
        auth_headers,
        sections=[
            {
                "type": "experience",
                "title": "Old Job",
                "start_date": "2010-01-01",
                "end_date": "2012-01-01",
            },
            {"type": "experience", "title": "Current Job", "start_date": "2020-01-01"},
            {
                "type": "experience",
                "title": "Middle Job",
                "start_date": "2013-01-01",
                "end_date": "2018-01-01",
            },
        ],
    )
    text = _pdf_text(_export(client, auth_headers, created["id"], "pdf").content)
    assert text.index("Current Job") < text.index("Middle Job") < text.index("Old Job")
    assert "Jan 2020 - Present" in _norm(text)


def test_pdf_special_characters_survive(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(
        auth_headers,
        full_name="Tom & Jerry <Admin>",
        summary="R&D lead: cut costs by >30% & latency <5ms. Used <b>bold</b> tags.",
        sections=[
            {
                "type": "experience",
                "title": "Dev & Ops <Lead>",
                "organization": "A&B Corp",
                "bullets": ["Improved x>y & y<z", "Handled <script>alert(1)</script> input"],
            }
        ],
    )
    response = _export(client, auth_headers, created["id"], "pdf")
    assert response.status_code == 200
    text = _norm(_pdf_text(response.content))
    assert "Tom & Jerry <Admin>" in text
    assert "R&D lead: cut costs by >30% & latency <5ms" in text
    assert "<b>bold</b>" in text  # markup is shown literally, not interpreted
    assert "Dev & Ops <Lead>" in text
    assert "A&B Corp" in text
    assert "x>y & y<z" in text
    assert "<script>alert(1)</script>" in text


def test_pdf_accents_quotes_and_symbols_export(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(
        auth_headers,
        title="Résumé 2026",
        full_name="José Müller",
        summary="Smart “quotes” and an em dash — plus emoji \U0001f600 and Zażó.",
    )
    response = _export(client, auth_headers, created["id"], "pdf")
    assert response.status_code == 200
    text = _pdf_text(response.content)
    assert "Jos" in text
    assert "quotes" in text
    disposition = response.headers["content-disposition"]
    assert disposition.isascii()


@pytest.mark.parametrize(
    ("field", "value", "expected"),
    [
        ("full_name", "Раджеш Кумар", "Р"),
        ("title", "Résumé 日本語", "日"),
        ("full_name", "Łukasz Nowak", "Ł"),
    ],
)
def test_pdf_refuses_letters_it_cannot_render(
    client: TestClient,
    auth_headers: Headers,
    make_resume: MakeResume,
    field: str,
    value: str,
    expected: str,
) -> None:
    """Never silently drop letters from a name: refuse and point to DOCX instead."""
    created = make_resume(auth_headers, **{field: value})
    response = _export(client, auth_headers, created["id"], "pdf")
    assert response.status_code == 400
    error = response.json()["error"]
    assert error["code"] == "BAD_REQUEST"
    assert expected in error["message"]
    assert "DOCX" in error["message"]


def test_docx_keeps_non_latin_text(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(auth_headers, full_name="Раджеш 田中 Łukasz")
    response = _export(client, auth_headers, created["id"], "docx")
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "Раджеш" in xml
    assert "田中" in xml
    assert "Łukasz" in xml


def test_pdf_empty_resume_still_exports(
    client: TestClient, auth_headers: Headers, template_id: int
) -> None:
    created = client.post(
        "/api/resumes", json={"template_id": template_id, "title": "Blank"}, headers=auth_headers
    ).json()
    response = _export(client, auth_headers, created["id"], "pdf")
    assert response.status_code == 200
    assert "Blank" in _pdf_text(response.content)


def test_pdf_long_resume_paginates(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    sections = [
        {
            "type": "experience",
            "title": f"Role {i}",
            "organization": "Org",
            "start_date": f"{2000 + i}-01-01",
            "bullets": [f"Achievement number {j} " + "word " * 30 for j in range(10)],
        }
        for i in range(20)
    ]
    created = make_resume(auth_headers, sections=sections)
    content = _export(client, auth_headers, created["id"], "pdf").content
    assert len(PdfReader(io.BytesIO(content)).pages) > 1


def test_pdf_has_no_images_or_annotations(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(auth_headers)
    reader = PdfReader(io.BytesIO(_export(client, auth_headers, created["id"], "pdf").content))
    for page in reader.pages:
        assert len(page.images) == 0
        assert "/Annots" not in page


def test_pdf_metadata_title(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(auth_headers)
    reader = PdfReader(io.BytesIO(_export(client, auth_headers, created["id"], "pdf").content))
    assert reader.metadata is not None
    assert reader.metadata.title == "Backend Engineer Resume"


def test_pdf_missing_resume_404(client: TestClient, auth_headers: Headers) -> None:
    response = _export(client, auth_headers, 9999, "pdf")
    assert response.status_code == 404
    assert response.headers["content-type"].startswith("application/json")


def test_export_other_users_resume_404(
    client: TestClient, auth_headers: Headers, other_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(auth_headers)
    assert _export(client, other_headers, created["id"], "pdf").status_code == 404
    assert _export(client, other_headers, created["id"], "docx").status_code == 404


def test_export_requires_auth(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(auth_headers)
    assert client.get(f"/api/resumes/{created['id']}/export/pdf").status_code == 401


def test_deactivated_template_still_exports(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume, db: Session
) -> None:
    created = make_resume(auth_headers)
    db.query(Template).filter(Template.id == created["template_id"]).update(
        {Template.is_active: False}
    )
    db.commit()
    assert _export(client, auth_headers, created["id"], "pdf").status_code == 200
    assert _export(client, auth_headers, created["id"], "docx").status_code == 200


@pytest.mark.parametrize("slug", ["classic", "modern-clean", "compact"])
def test_each_seeded_template_exports(
    client: TestClient, auth_headers: Headers, db: Session, slug: str
) -> None:
    tid = db.query(Template).filter(Template.slug == slug).one().id
    created = client.post(
        "/api/resumes", json=resume_payload(int(tid)), headers=auth_headers
    ).json()
    assert _export(client, auth_headers, created["id"], "pdf").status_code == 200
    assert _export(client, auth_headers, created["id"], "docx").status_code == 200


def test_export_does_not_mutate_resume(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(auth_headers)
    _export(client, auth_headers, created["id"], "pdf")
    _export(client, auth_headers, created["id"], "docx")
    assert client.get(f"/api/resumes/{created['id']}", headers=auth_headers).json() == created


# -------------------------------------------------------------------- DOCX
def test_docx_export_headers_and_zip(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(auth_headers)
    response = _export(client, auth_headers, created["id"], "docx")
    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX_TYPE
    assert (
        response.headers["content-disposition"]
        == 'attachment; filename="Backend_Engineer_Resume.docx"'
    )
    assert response.headers["cache-control"] == "no-store"
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert archive.testzip() is None
        names = archive.namelist()
        assert "word/document.xml" in names
        assert not any(n.startswith("word/media/") for n in names)


def test_docx_structure_is_ats_safe(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(auth_headers)
    content = _export(client, auth_headers, created["id"], "docx").content
    document = Document(io.BytesIO(content))
    assert document.tables == []
    assert len(document.inline_shapes) == 0
    section = document.sections[0]
    assert not any(p.text.strip() for p in section.header.paragraphs)
    assert not any(p.text.strip() for p in section.footer.paragraphs)
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "<w:tbl>" not in xml
    assert 'w:num="' not in xml  # no multi-column layout
    assert "<w:txbxContent" not in xml

    texts = [p.text for p in document.paragraphs]
    headings = [
        p.text for p in document.paragraphs if p.style is not None and p.style.name == "Heading 1"
    ]
    assert headings == ["Summary", "Experience", "Education", "Skills"]
    assert texts[0] == "Alice Example"
    assert "alice@example.com" in texts[1]
    bullets = [
        p.text for p in document.paragraphs if p.style is not None and p.style.name == "List Bullet"
    ]
    assert bullets == ["Built the billing API", "Cut p99 latency by 40%"]
    assert document.core_properties.title == "Backend Engineer Resume"


def test_docx_special_characters(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(auth_headers, full_name="Tom & Jerry <Admin>", summary="a>b & c<d")
    content = _export(client, auth_headers, created["id"], "docx").content
    texts = [p.text for p in Document(io.BytesIO(content)).paragraphs]
    assert "Tom & Jerry <Admin>" in texts
    assert "a>b & c<d" in texts


# -------------------------------------------------------- service helpers
@pytest.mark.parametrize(
    ("title", "expected"),
    [
        ("My Resume", "My_Resume.pdf"),
        ("  spaced   out  ", "spaced_out.pdf"),
        ("../../etc/passwd", "etc_passwd.pdf"),
        ('a"b\r\nX-Evil: 1', "a_b_X_Evil_1.pdf"),
        ("日本語", "resume.pdf"),
        ("", "resume.pdf"),
        ("x" * 200, "x" * 80 + ".pdf"),
    ],
)
def test_safe_filename(title: str, expected: str) -> None:
    assert export_service.safe_filename(title, "pdf") == expected


@pytest.mark.parametrize(
    ("font", "regular"),
    [
        ("Times New Roman", "Times-Roman"),
        ("Georgia", "Times-Roman"),
        ("Courier New", "Courier"),
        ("Arial", "Helvetica"),
        ("Calibri", "Helvetica"),
        ("Sans Serif", "Helvetica"),
        ("", "Helvetica"),
        ("Papyrus", "Helvetica"),
    ],
)
def test_pdf_font_mapping(font: str, regular: str) -> None:
    assert export_service.pdf_fonts_for(font)[0] == regular


def test_invalid_accent_color_falls_back_to_black() -> None:
    doc = export_service.ExportDocument(
        title="T",
        name="N",
        contact=[],
        sections=[export_service.ExportSection("Summary", paragraph="x")],
        font_family="Arial",
        accent_color=export_service._accent("not-a-colour"),
    )
    assert doc.accent_color == "#000000"
    assert export_service.render_pdf(doc).startswith(b"%PDF")


def test_clean_strips_control_characters_and_collapses_newlines() -> None:
    assert export_service._clean("a\x00b\x07c") == "abc"
    assert export_service._clean("line1\r\nline2\tx") == "line1 line2 x"
    assert export_service._clean(None) == ""


def test_build_document_dedupes_skills_and_skips_untitled(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume, db: Session
) -> None:
    from app.models import Resume

    created = make_resume(auth_headers, skills=["Python", "python", "Go"])
    resume = db.get(Resume, created["id"])
    template = db.get(Template, created["template_id"])
    assert resume is not None and template is not None
    document = export_service.build_export_document(resume, template)
    skills = next(s for s in document.sections if s.heading == "Skills")
    assert skills.paragraph == "Python, Go"
    assert export_service.render_resume_pdf(resume, template).startswith(b"%PDF")
    assert export_service.render_resume_docx(resume, template)[:2] == b"PK"


def test_end_date_only_is_rendered_without_range(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(
        auth_headers,
        sections=[{"type": "education", "title": "MSc", "end_date": "2018-06-01"}],
    )
    text = _norm(_pdf_text(_export(client, auth_headers, created["id"], "pdf").content))
    assert "Jun 2018" in text
    assert "Present" not in text
