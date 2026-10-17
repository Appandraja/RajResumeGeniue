"""Multi-format export: media types, per-format content, Unicode, escaping, routing."""

import io
import re
import zipfile
from collections.abc import Callable
from html.parser import HTMLParser
from typing import Any

import pytest
from docx import Document
from fastapi.testclient import TestClient
from odf import teletype
from odf.opendocument import load as load_odt
from odf.text import H, List, ListItem
from odf.text import P as OdfP
from pypdf import PdfReader
from striprtf.striprtf import rtf_to_text

from app.config import settings
from app.rate_limit import limit_export
from app.services import export_service as svc

Headers = dict[str, str]
MakeResume = Callable[..., dict[str, Any]]

FORMATS: dict[str, tuple[str, str]] = {
    "pdf": ("application/pdf", "pdf"),
    "docx": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "docx",
    ),
    "odt": ("application/vnd.oasis.opendocument.text", "odt"),
    "rtf": ("application/rtf", "rtf"),
    "txt": ("text/plain; charset=utf-8", "txt"),
    "md": ("text/markdown; charset=utf-8", "md"),
    "html": ("text/html; charset=utf-8", "html"),
}
UNICODE_FORMATS = [f for f in FORMATS if f != "pdf"]

SECTIONS = [
    {
        "type": "experience",
        "title": "Old Job",
        "organization": "OldCo",
        "start_date": "2010-01-01",
        "end_date": "2012-01-01",
        "bullets": ["Did old things"],
    },
    {
        "type": "experience",
        "title": "Current Job",
        "organization": "NewCo",
        "location": "Remote",
        "start_date": "2020-01-01",
        "bullets": ["Shipped the platform", "Mentored four engineers"],
    },
    {
        "type": "education",
        "title": "BSc Computer Science",
        "organization": "State University",
        "start_date": "2012-09-01",
        "end_date": "2016-06-01",
    },
]


def _get(client: TestClient, headers: Headers, rid: int, fmt: str) -> Any:
    return client.get(f"/api/resumes/{rid}/export/{fmt}", headers=headers)


def _norm(text: str) -> str:
    return " ".join(text.split())


def _pdf_text(content: bytes) -> str:
    return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(content)).pages)


def _docx_text(content: bytes) -> str:
    return "\n".join(p.text for p in Document(io.BytesIO(content)).paragraphs)


def _odt_text(content: bytes) -> str:
    doc = load_odt(io.BytesIO(content))
    return "\n".join(teletype.extractText(el) for el in doc.text.childNodes)


def _rtf_text(content: bytes) -> str:
    return str(rtf_to_text(content.decode("ascii")))


def _extract(fmt: str, content: bytes) -> str:
    if fmt == "pdf":
        return _pdf_text(content)
    if fmt == "docx":
        return _docx_text(content)
    if fmt == "odt":
        return _odt_text(content)
    if fmt == "rtf":
        return _rtf_text(content)
    return content.decode("utf-8")


def _strip_md(text: str) -> str:
    """Undo Markdown backslash escapes so tests can compare against plain text."""
    return re.sub(r"\\([\\`*_{}\[\]()#+\-.!|~<>&])", r"\1", text)


def _plain(fmt: str, content: bytes) -> str:
    text = _extract(fmt, content)
    if fmt == "md":
        text = _strip_md(text)
    return text


@pytest.fixture
def full_resume(client: TestClient, auth_headers: Headers, make_resume: MakeResume) -> int:
    return int(make_resume(auth_headers, sections=SECTIONS)["id"])


# ------------------------------------------------------------ media types
@pytest.mark.parametrize("fmt", list(FORMATS))
def test_every_format_returns_200_with_headers(
    client: TestClient, auth_headers: Headers, full_resume: int, fmt: str
) -> None:
    media_type, ext = FORMATS[fmt]
    response = _get(client, auth_headers, full_resume, fmt)
    assert response.status_code == 200
    assert response.headers["content-type"] == media_type
    assert response.headers["cache-control"] == "no-store"
    disposition = response.headers["content-disposition"]
    assert disposition == f'attachment; filename="Backend_Engineer_Resume.{ext}"'
    assert disposition.isascii()
    assert len(response.content) > 100


def test_registry_matches_supported_list() -> None:
    assert set(svc.EXPORT_FORMATS) == set(FORMATS)
    for key, (media_type, _) in FORMATS.items():
        assert svc.EXPORT_FORMATS[key].media_type == media_type


# ----------------------------------------------------------- content/order
@pytest.mark.parametrize("fmt", list(FORMATS))
def test_content_present_and_in_order(
    client: TestClient, auth_headers: Headers, full_resume: int, fmt: str
) -> None:
    text = _norm(_plain(fmt, _get(client, auth_headers, full_resume, fmt).content))
    ordered = [
        "Alice Example",
        "alice@example.com",
        "Berlin, Germany",
        "Backend engineer with a decade",
        "Current Job",
        "Shipped the platform",
        "Mentored four engineers",
        "Old Job",
        "BSc Computer Science",
        "State University",
        "Python, FastAPI, PostgreSQL, Docker, AWS",
    ]
    positions = []
    for expected in ordered:
        assert expected in text, (fmt, expected)
        positions.append(text.index(expected))
    assert positions == sorted(positions)
    heading_pos = [text.upper().index(h.upper()) for h in ("Summary", "Experience", "Education")]
    assert heading_pos == sorted(heading_pos)
    assert text.upper().index("SKILLS") > text.index("BSc Computer Science")


@pytest.mark.parametrize("fmt", list(FORMATS))
def test_present_for_open_end_date_and_ranges(
    client: TestClient, auth_headers: Headers, full_resume: int, fmt: str
) -> None:
    text = _norm(_plain(fmt, _get(client, auth_headers, full_resume, fmt).content))
    assert "Jan 2020 - Present" in text
    assert "Jan 2010 - Jan 2012" in text
    assert "Sep 2012 - Jun 2016" in text


@pytest.mark.parametrize("fmt", list(FORMATS))
def test_empty_optional_fields_make_no_blank_headings(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume, fmt: str
) -> None:
    created = make_resume(
        auth_headers,
        summary="",
        skills=[],
        phone="",
        location="",
        linkedin_url="",
        sections=[{"type": "experience", "title": "Engineer", "start_date": "2020-01-01"}],
    )
    text = _plain(fmt, _get(client, auth_headers, created["id"], fmt).content)
    upper = text.upper()
    assert "EXPERIENCE" in upper
    for heading in ("SUMMARY", "SKILLS", "EDUCATION", "PROJECTS", "CERTIFICATIONS"):
        assert heading not in upper, (fmt, heading)


@pytest.mark.parametrize("fmt", list(FORMATS))
def test_blank_resume_still_exports_title(
    client: TestClient, auth_headers: Headers, template_id: int, fmt: str
) -> None:
    created = client.post(
        "/api/resumes",
        json={"template_id": template_id, "title": "Blank Doc"},
        headers=auth_headers,
    ).json()
    response = _get(client, auth_headers, created["id"], fmt)
    assert response.status_code == 200
    assert "Blank Doc" in _plain(fmt, response.content)


# ----------------------------------------------------------------- Unicode
UNICODE_NAME = "Раджеш 田中 مرحبا Łukasz José"


@pytest.mark.parametrize("fmt", UNICODE_FORMATS)
def test_unicode_survives(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume, fmt: str
) -> None:
    created = make_resume(
        auth_headers,
        full_name=UNICODE_NAME,
        summary="Ingénieur à Zürich — 日本語の経験 и русский текст",
        sections=[
            {
                "type": "experience",
                "title": "Разработчик",
                "organization": "株式会社",
                "start_date": "2020-01-01",
                "bullets": ["Построил систему ∑ café"],
            }
        ],
    )
    response = _get(client, auth_headers, created["id"], fmt)
    assert response.status_code == 200
    text = _plain(fmt, response.content)
    for expected in (
        UNICODE_NAME,
        "Ingénieur à Zürich",
        "日本語の経験",
        "Разработчик",
        "株式会社",
        "Построил систему ∑ café",
    ):
        assert expected in text, (fmt, expected)


def test_pdf_unicode_refused_with_pointer_to_other_formats(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(auth_headers, full_name=UNICODE_NAME)
    response = _get(client, auth_headers, created["id"], "pdf")
    assert response.status_code == 400
    message = response.json()["error"]["message"]
    assert "Please download DOCX, ODT, RTF, TXT, HTML or Markdown instead" in message


@pytest.mark.parametrize("fmt", ["rtf", "txt", "md", "html", "odt"])
def test_astral_characters_do_not_break_rendering(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume, fmt: str
) -> None:
    created = make_resume(auth_headers, summary="Launch \U0001f680 done")
    assert _get(client, auth_headers, created["id"], fmt).status_code == 200


def test_rtf_astral_uses_utf16_surrogate_pair() -> None:
    assert svc._rtf("\U0001f680") == "\\u-10179?\\u-8576?"
    assert svc._rtf("é") == "\\u233?"
    assert svc._rtf("\u4e2d") == "\\u20013?"


# ------------------------------------------------------------- per format
def test_rtf_escapes_special_characters(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(
        auth_headers,
        full_name="A{b}c\\d",
        summary="path C:\\temp {group} \\par injected",
    )
    content = _get(client, auth_headers, created["id"], "rtf").content
    raw = content.decode("ascii")
    assert "A\\{b\\}c\\\\d" in raw
    assert "C:\\\\temp \\{group\\} \\\\par injected" in raw
    assert raw.count("{") - raw.count("\\{") == raw.count("}") - raw.count("\\}")
    assert raw.startswith("{\\rtf1")
    text = _rtf_text(content)
    assert "A{b}c\\d" in text
    assert "C:\\temp {group} \\par injected" in text
    assert "\\tbl" not in raw and "\\trowd" not in raw


def test_txt_structure(client: TestClient, auth_headers: Headers, full_resume: int) -> None:
    response = _get(client, auth_headers, full_resume, "txt")
    body = response.content.decode("utf-8")
    assert not response.content.startswith(b"\xef\xbb\xbf")
    lines = body.splitlines()
    assert lines[0] == "Alice Example"
    assert {"SUMMARY", "EXPERIENCE", "EDUCATION", "SKILLS"} <= set(lines)
    assert "- Shipped the platform" in lines
    assert "- Mentored four engineers" in lines
    assert all(len(line) <= 80 for line in lines)
    assert "\t" not in body


def test_txt_wraps_long_bullets_with_hanging_indent(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    long_bullet = " ".join(f"word{i}" for i in range(60))
    created = make_resume(
        auth_headers,
        sections=[
            {
                "type": "experience",
                "title": "Eng",
                "start_date": "2020-01-01",
                "bullets": [long_bullet],
            }
        ],
    )
    lines = _get(client, auth_headers, created["id"], "txt").content.decode().splitlines()
    idx = next(i for i, line in enumerate(lines) if line.startswith("- word0"))
    assert lines[idx + 1].startswith("  word")
    assert all(len(line) <= 80 for line in lines)


def test_md_structure_and_escaping(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(
        auth_headers,
        full_name="Ann *Star* [Dev]",
        summary="Uses `code`, # not heading, <b>tags</b> and a_b_c | pipes",
        sections=[
            {
                "type": "experience",
                "title": "- Lead",
                "start_date": "2020-01-01",
                "bullets": ["# fake heading", "1. fake list", "![img](http://x/y.png)"],
            }
        ],
    )
    body = _get(client, auth_headers, created["id"], "md").content.decode("utf-8")
    lines = body.splitlines()
    assert lines[0] == "# Ann \\*Star\\* \\[Dev\\]"
    assert "## Summary" in lines and "## Experience" in lines and "## Skills" in lines
    assert "\\`code\\`" in body
    assert "\\# not heading" in body
    assert "\\<b\\>tags\\</b\\>" in body
    assert "a\\_b\\_c \\| pipes" in body
    assert "### \\- Lead" in lines
    assert "- \\# fake heading" in lines
    assert "- 1\\. fake list" in lines
    assert "- !\\[img\\](http://x/y.png)" in lines
    assert not any(line.startswith("# ") for line in lines[1:])
    assert "- Python" not in body  # skills stay a single comma-separated paragraph


class _Collect(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.tags: list[str] = []
        self.stack: list[str] = []
        self.errors: list[str] = []
        self.attrs: list[tuple[str, str]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self.tags.append(tag)
        self.attrs.extend((tag, k) for k, _ in attrs)
        if tag not in {"meta", "link", "br", "hr", "img", "input"}:
            self.stack.append(tag)

    def handle_endtag(self, tag: str) -> None:
        if not self.stack or self.stack.pop() != tag:
            self.errors.append(tag)


def test_html_is_valid_semantic_and_script_free(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(
        auth_headers,
        full_name='Tom & "Jerry" <Admin>',
        summary="</p><script>alert(1)</script>",
        sections=[
            {
                "type": "experience",
                "title": "<img src=x onerror=alert(1)>",
                "organization": "A&B",
                "start_date": "2020-01-01",
                "bullets": ["Handled <script>alert(2)</script> input"],
            }
        ],
    )
    response = _get(client, auth_headers, created["id"], "html")
    body = response.content.decode("utf-8")
    assert body.startswith("<!DOCTYPE html>")
    assert '<meta charset="utf-8">' in body
    assert "<script" not in body.lower()
    assert "<img" not in body.lower()
    assert "&lt;script&gt;alert(2)&lt;/script&gt;" in body
    assert "Tom &amp; &quot;Jerry&quot; &lt;Admin&gt;" in body
    assert "A&amp;B" in body
    for external in ("<link", "@import", "url("):  # no external resources
        assert external not in body
    assert "<style>" in body
    parser = _Collect()
    parser.feed(body)
    assert parser.errors == []
    assert parser.stack == []
    for tag in ("h1", "h2", "h3", "ul", "li", "title"):
        assert tag in parser.tags
    assert parser.tags.count("h1") == 1
    assert "table" not in parser.tags and "img" not in parser.tags
    assert not any(attr.startswith("on") or attr in {"src", "href"} for _, attr in parser.attrs)


def test_html_uses_title_as_h1_when_name_missing(
    client: TestClient, auth_headers: Headers, template_id: int
) -> None:
    created = client.post(
        "/api/resumes",
        json={"template_id": template_id, "title": "Only Title"},
        headers=auth_headers,
    ).json()
    body = _get(client, auth_headers, created["id"], "html").content.decode()
    assert "<h1>Only Title</h1>" in body


def test_odt_structure_is_ats_safe(
    client: TestClient, auth_headers: Headers, full_resume: int
) -> None:
    content = _get(client, auth_headers, full_resume, "odt").content
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        assert archive.testzip() is None
        names = archive.namelist()
        assert names[0] == "mimetype"
        assert archive.read("mimetype") == b"application/vnd.oasis.opendocument.text"
        assert not any(n.startswith("Pictures/") for n in names)
        xml = archive.read("content.xml").decode("utf-8")
    assert "<table:table" not in xml
    assert "<draw:" not in xml
    doc = load_odt(io.BytesIO(content))
    headings = [teletype.extractText(h) for h in doc.getElementsByType(H)]
    assert headings == ["Summary", "Experience", "Education", "Skills"]
    items = [
        teletype.extractText(p)
        for li in doc.getElementsByType(ListItem)
        for p in li.getElementsByType(OdfP)
    ]
    assert items == ["Shipped the platform", "Mentored four engineers", "Did old things"]
    assert len(doc.getElementsByType(List)) == 2


def test_odt_special_characters(
    client: TestClient, auth_headers: Headers, make_resume: MakeResume
) -> None:
    created = make_resume(auth_headers, full_name="Tom & Jerry <Admin>", summary="a>b & c<d")
    text = _odt_text(_get(client, auth_headers, created["id"], "odt").content)
    assert "Tom & Jerry <Admin>" in text
    assert "a>b & c<d" in text


def test_font_name_sanitised_in_css_and_rtf() -> None:
    doc = svc.ExportDocument(
        title="T",
        name="N",
        contact=[],
        sections=[],
        font_family='Evil"; } body { background: url(x) ',
        accent_color="#123456",
    )
    html_body = svc.render_html(doc).decode()
    assert "url(" not in html_body
    assert '"Evil body background urlx"' in html_body
    assert svc._font_name(doc.font_family) == "Evil body background urlx"
    assert b"Evil body background urlx;" in svc.render_rtf(doc)
    assert svc._font_name("") == "Arial"


# -------------------------------------------------------------- routing
@pytest.mark.parametrize("fmt", ["exe", "PDF", "doc", "pdf.zip", "xml"])
def test_unknown_format_is_400_listing_supported(
    client: TestClient, auth_headers: Headers, full_resume: int, fmt: str
) -> None:
    response = _get(client, auth_headers, full_resume, fmt)
    assert response.status_code == 400
    error = response.json()["error"]
    assert error["code"] == "BAD_REQUEST"
    for supported in FORMATS:
        assert supported in error["message"]


@pytest.mark.parametrize("fmt", list(FORMATS))
def test_requires_auth(client: TestClient, full_resume: int, fmt: str) -> None:
    assert client.get(f"/api/resumes/{full_resume}/export/{fmt}").status_code == 401


@pytest.mark.parametrize("fmt", list(FORMATS))
def test_other_users_resume_is_404(
    client: TestClient, other_headers: Headers, full_resume: int, fmt: str
) -> None:
    assert _get(client, other_headers, full_resume, fmt).status_code == 404


def test_missing_resume_is_404(client: TestClient, auth_headers: Headers) -> None:
    assert _get(client, auth_headers, 99999, "txt").status_code == 404


def test_rate_limit_applies_to_every_format(
    client: TestClient,
    auth_headers: Headers,
    full_resume: int,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "RATE_LIMIT_ENABLED", True)
    limit_export.reset()
    formats = list(FORMATS)
    for i in range(limit_export.limit):
        assert _get(client, auth_headers, full_resume, formats[i % len(formats)]).status_code == 200
    blocked = _get(client, auth_headers, full_resume, "md")
    assert blocked.status_code == 429
    limit_export.reset()


def test_old_explicit_paths_still_resolve(
    client: TestClient, auth_headers: Headers, full_resume: int
) -> None:
    assert _get(client, auth_headers, full_resume, "pdf").content.startswith(b"%PDF")
    assert _get(client, auth_headers, full_resume, "docx").content[:2] == b"PK"


def test_safe_filename_folds_accents() -> None:
    assert svc.safe_filename("Ünï Resume", "md") == "Uni_Resume.md"
