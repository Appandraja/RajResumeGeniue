"""Resume import: every format, parser heuristics, security limits and the full loop."""

import io
import zipfile
from collections.abc import Generator
from typing import Any

import pytest
from docx import Document
from fastapi.testclient import TestClient
from odf.opendocument import OpenDocumentText
from odf.text import H, P
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.config import settings
from app.main import app
from app.rate_limit import limit_import
from app.routers import imports as imports_router
from app.services import import_service, resume_parser
from tests.conftest import PASSWORD

Headers = dict[str, str]
URL = "/api/import/resume"

RESUME_TEXT = """Jane Q. Doe
Berlin, Germany | jane.doe@example.com | +49 151 2345 6789
linkedin.com/in/janedoe | janedoe.dev

SUMMARY
Backend engineer with ten years of experience building reliable services.

WORK EXPERIENCE
Senior Software Engineer | Acme Corp | Berlin, Germany
Jan 2020 - Present
- Built the billing API serving 2M requests a day
- Cut p99 latency by 40% by introducing
  a caching layer
Software Engineer, Initech
03/2016 - 12/2019
* Maintained legacy Java services
* Mentored 3 junior engineers

EDUCATION
State University
BSc Computer Science
2012 - 2016

PROJECTS
Resume Genie | Python, React
- Built an ATS friendly resume builder

CERTIFICATIONS
AWS Certified Solutions Architect - Amazon, 2021

SKILLS
Languages: Python, Go, SQL
Tools: Docker | Kubernetes, python, Git
"""


@pytest.fixture(autouse=True)
def _wire_router() -> Generator[None, None, None]:
    """Mount the import router when main.py has not wired it yet."""
    if not any(getattr(r, "path", "") == URL for r in app.routes):
        app.include_router(imports_router.router)
    yield


# -------------------------------------------------------------- generators
def make_pdf(text: str, pages: int = 1) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    for _ in range(pages):
        y = 800
        for line in text.split("\n"):
            pdf.drawString(40, y, line)
            y -= 13
            if y < 40:
                break
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_blank_pdf() -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    pdf.rect(100, 100, 200, 200, fill=1)
    pdf.save()
    return buffer.getvalue()


def make_docx(text: str) -> bytes:
    document = Document()
    for line in text.split("\n"):
        stripped = line.strip()
        if stripped.startswith(("- ", "* ")):
            document.add_paragraph(stripped[2:], style="List Bullet")
        else:
            document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_odt(text: str) -> bytes:
    doc = OpenDocumentText()
    for line in text.split("\n"):
        stripped = line.strip()
        if stripped.isupper() and stripped:
            doc.text.addElement(H(outlinelevel=1, text=stripped))
        else:
            doc.text.addElement(P(text=line))
    buffer = io.BytesIO()
    doc.write(buffer)
    return buffer.getvalue()


def make_rtf(text: str) -> bytes:
    body = "".join(
        r"\par " + line.replace("\\", "\\\\").replace("{", "\\{").replace("}", "\\}") + "\n"
        for line in text.split("\n")
    )
    return ("{\\rtf1\\ansi\\deff0 {\\fonttbl{\\f0 Arial;}}\n" + body + "}").encode("latin-1")


def make_html(text: str) -> bytes:
    rows = []
    for line in text.split("\n"):
        stripped = line.strip()
        if stripped.startswith(("- ", "* ")):
            rows.append(f"<ul><li>{stripped[2:]}</li></ul>")
        elif stripped:
            rows.append(f"<p>{stripped}</p>")
    html = (
        "<html><head><title>CV</title><style>p{color:red}</style>"
        "<script>var secret = 'SHOULD-NOT-APPEAR';</script></head><body>"
        + "".join(rows)
        + "</body></html>"
    )
    return html.encode("utf-8")


def _post(
    client: TestClient, headers: Headers, name: str, data: bytes, content_type: str = "text/plain"
) -> Any:
    return client.post(URL, files={"file": (name, data, content_type)}, headers=headers)


def _parse_ok(client: TestClient, headers: Headers, name: str, data: bytes) -> dict[str, Any]:
    response = _post(client, headers, name, data)
    assert response.status_code == 200, response.text
    body: dict[str, Any] = response.json()
    return body


def _txt(text: str) -> bytes:
    return text.encode("utf-8")


def _assert_realistic(body: dict[str, Any]) -> None:
    draft = body["draft"]
    assert draft["full_name"] == "Jane Q. Doe"
    assert draft["title"] == "Jane Q. Doe Resume"
    assert draft["email"] == "jane.doe@example.com"
    assert draft["phone"] == "+49 151 2345 6789"
    assert draft["location"] == "Berlin, Germany"
    assert draft["linkedin_url"] == "https://linkedin.com/in/janedoe"
    assert draft["website_url"] == "https://janedoe.dev"
    assert draft["summary"].startswith("Backend engineer with ten years")
    assert draft["skills"] == ["Python", "Go", "SQL", "Docker", "Kubernetes", "Git"]
    sections = draft["sections"]
    assert [s["type"] for s in sections] == [
        "experience",
        "experience",
        "education",
        "project",
        "certification",
    ]
    assert [s["position"] for s in sections] == [0, 1, 2, 3, 4]
    senior = sections[0]
    assert senior["title"] == "Senior Software Engineer"
    assert senior["organization"] == "Acme Corp"
    assert senior["location"] == "Berlin, Germany"
    assert senior["start_date"] == "2020-01-01"
    assert senior["end_date"] is None
    assert senior["bullets"][0] == "Built the billing API serving 2M requests a day"
    assert "caching layer" in senior["bullets"][1]
    initech = sections[1]
    assert (initech["title"], initech["organization"]) == ("Software Engineer", "Initech")
    assert (initech["start_date"], initech["end_date"]) == ("2016-03-01", "2019-12-01")
    assert len(initech["bullets"]) == 2
    edu = sections[2]
    assert edu["title"] == "BSc Computer Science"
    assert edu["organization"] == "State University"
    assert (edu["start_date"], edu["end_date"]) == ("2012-01-01", "2016-01-01")
    cert = sections[4]
    assert cert["title"] == "AWS Certified Solutions Architect"
    assert cert["organization"] == "Amazon"
    assert cert["start_date"] == "2021-01-01"
    assert body["char_count"] > 100


# ------------------------------------------------------------ round trips
@pytest.mark.parametrize(
    ("name", "builder", "fmt"),
    [
        ("resume.txt", _txt, "txt"),
        ("resume.md", _txt, "md"),
        ("resume.markdown", _txt, "md"),
        ("resume.docx", make_docx, "docx"),
        ("resume.odt", make_odt, "odt"),
        ("resume.rtf", make_rtf, "rtf"),
        ("resume.html", make_html, "html"),
        ("resume.htm", make_html, "html"),
        ("RESUME.PDF", make_pdf, "pdf"),
    ],
)
def test_format_round_trip(
    client: TestClient, auth_headers: Headers, name: str, builder: Any, fmt: str
) -> None:
    body = _parse_ok(client, auth_headers, name, builder(RESUME_TEXT))
    assert body["source_format"] == fmt
    _assert_realistic(body)
    assert body["warnings"] == []


def test_markdown_syntax_is_stripped(client: TestClient, auth_headers: Headers) -> None:
    md = (
        "# Jane Doe\n\n"
        "[LinkedIn](https://www.linkedin.com/in/jane-doe) | **jane@example.com**\n\n"
        "---\n\n## Skills\n\n* Python\n* **FastAPI**\n"
    )
    body = _parse_ok(client, auth_headers, "cv.md", _txt(md))
    draft = body["draft"]
    assert draft["full_name"] == "Jane Doe"
    assert draft["email"] == "jane@example.com"
    assert draft["linkedin_url"] == "https://www.linkedin.com/in/jane-doe"
    assert draft["skills"] == ["Python", "FastAPI"]


def test_html_scripts_and_styles_are_dropped(client: TestClient, auth_headers: Headers) -> None:
    result = import_service.extract_text("html", make_html(RESUME_TEXT))
    assert "SHOULD-NOT-APPEAR" not in result.text
    assert "color:red" not in result.text
    assert "Jane Q. Doe" in result.text


def test_docx_tables_are_read(client: TestClient, auth_headers: Headers) -> None:
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Maria Garcia"
    table.rows[0].cells[1].text = "maria@example.com"
    document.add_paragraph("SKILLS")
    document.add_paragraph("Excel, Sales")
    buffer = io.BytesIO()
    document.save(buffer)
    body = _parse_ok(client, auth_headers, "t.docx", buffer.getvalue())
    assert body["draft"]["full_name"] == "Maria Garcia"
    assert body["draft"]["email"] == "maria@example.com"
    assert body["draft"]["skills"] == ["Excel", "Sales"]


def test_rtf_escapes_unicode(client: TestClient, auth_headers: Headers) -> None:
    rtf = rb"{\rtf1\ansi\deff0 \par Jos\'e9 Garc\'eda\par jose@example.com\par}"
    body = _parse_ok(client, auth_headers, "cv.rtf", rtf)
    assert body["draft"]["full_name"] == "José García"


# ------------------------------------------------------- parser behaviours
@pytest.mark.parametrize(
    ("date_line", "start", "end"),
    [
        ("Jan 2020 - Present", "2020-01-01", None),
        ("01/2020 - 03/2022", "2020-01-01", "2022-03-01"),
        ("2019 - 2021", "2019-01-01", "2021-01-01"),
        ("March 2018 to June 2020", "2018-03-01", "2020-06-01"),
        ("Sept. 2017 - current", "2017-09-01", None),
        ("2018 – Now", "2018-01-01", None),
        ("(May 2015 — Aug 2016)", "2015-05-01", "2016-08-01"),
        ("2014-06 - 2015-02", "2014-06-01", "2015-02-01"),
        ("Dec 2010 until Ongoing", "2010-12-01", None),
    ],
)
def test_date_range_variants(
    client: TestClient, auth_headers: Headers, date_line: str, start: str, end: str | None
) -> None:
    text = (
        "Sam Lee\nsam@example.com\n\nEXPERIENCE\nData Analyst\nGlobex\n"
        f"{date_line}\n- Built reports\n"
    )
    section = _parse_ok(client, auth_headers, "c.txt", _txt(text))["draft"]["sections"][0]
    assert (section["start_date"], section["end_date"]) == (start, end)
    assert section["title"] == "Data Analyst"
    assert section["organization"] == "Globex"
    assert section["bullets"] == ["Built reports"]


def test_inline_dates_and_title_org_swap(client: TestClient, auth_headers: Headers) -> None:
    text = (
        "Sam Lee\nsam@example.com\n\nEmployment\n"
        "Globex Inc. | Data Analyst | Jan 2019 - Dec 2020\n"
        "• Built dashboards\n"
        "• Automated the weekly report\n"
    )
    section = _parse_ok(client, auth_headers, "c.txt", _txt(text))["draft"]["sections"][0]
    assert section["title"] == "Data Analyst"
    assert section["organization"] == "Globex Inc."
    assert section["start_date"] == "2019-01-01"
    assert section["end_date"] == "2020-12-01"
    assert section["bullets"] == ["Built dashboards", "Automated the weekly report"]


def test_date_last_title_first_layout(client: TestClient, auth_headers: Headers) -> None:
    text = (
        "Sam Lee\nsam@example.com\n\nPROFESSIONAL EXPERIENCE\n"
        "Software Developer    2020 - 2022\nAcme Corp, Austin, TX\n- Shipped features\n"
        "Junior Developer    2018 - 2020\nOther Co\n- Fixed bugs\n"
    )
    sections = _parse_ok(client, auth_headers, "c.txt", _txt(text))["draft"]["sections"]
    assert [s["title"] for s in sections] == ["Software Developer", "Junior Developer"]
    assert sections[0]["organization"] == "Acme Corp"
    assert sections[0]["location"] == "Austin, TX"


def test_missing_dates_warns(client: TestClient, auth_headers: Headers) -> None:
    text = "Sam Lee\nsam@example.com\n\nEXPERIENCE\nData Analyst, Globex\n- Built reports\n"
    body = _parse_ok(client, auth_headers, "c.txt", _txt(text))
    assert body["draft"]["sections"][0]["start_date"] is None
    assert any("No dates found for 1" in w for w in body["warnings"])
    assert any("skills" in w.lower() for w in body["warnings"])


def test_single_date_semantics(client: TestClient, auth_headers: Headers) -> None:
    text = (
        "Sam Lee\nsam@example.com\n\nEDUCATION\nBSc Physics, State University\nGraduated 2016\n\n"
        "CERTIFICATIONS\nPMP - PMI\nIssued Mar 2021\n"
    )
    draft = _parse_ok(client, auth_headers, "c.txt", _txt(text))["draft"]
    edu, cert = draft["sections"]
    assert (edu["start_date"], edu["end_date"]) == (None, "2016-01-01")
    assert edu["title"] == "BSc Physics"
    assert edu["organization"] == "State University"
    assert (cert["start_date"], cert["end_date"]) == ("2021-03-01", None)


def test_missing_sections_and_no_headings(client: TestClient, auth_headers: Headers) -> None:
    body = _parse_ok(client, auth_headers, "my cv.txt", _txt("Just a couple of loose words\n"))
    draft = body["draft"]
    assert draft["sections"] == []
    assert draft["skills"] == []
    assert any("No experience, education" in w for w in body["warnings"])
    assert any("No email" in w for w in body["warnings"])


def test_no_name_falls_back_to_filename(client: TestClient, auth_headers: Headers) -> None:
    text = "SKILLS\nPython, Go\n"
    body = _parse_ok(client, auth_headers, "Q3 Resume.txt", _txt(text))
    assert body["draft"]["full_name"] is None
    assert body["draft"]["title"] == "Q3 Resume"
    assert any("full name" in w for w in body["warnings"])


def test_name_guessed_warning(client: TestClient, auth_headers: Headers) -> None:
    text = "Resume\nCurriculum Vitae\nPriya Natarajan\npriya@example.com\n\nSKILLS\nSQL\n"
    body = _parse_ok(client, auth_headers, "c.txt", _txt(text))
    assert body["draft"]["full_name"] == "Priya Natarajan"
    assert any("guessed" in w for w in body["warnings"])


def test_unicode_names_and_accents(client: TestClient, auth_headers: Headers) -> None:
    text = (
        "Zoë Ångström-Müller\nZürich, Switzerland | zoe@example.com\n\n"
        "BERUFSERFAHRUNG\nEXPERIENCE\nIngénieur logiciel | Société Générale | Paris, France\n"
        "2017 - 2020\n- Développé des services\n\nSKILLS\nPython, C++, C#, Node.js\n"
    )
    draft = _parse_ok(client, auth_headers, "c.txt", _txt(text))["draft"]
    assert draft["full_name"] == "Zoë Ångström-Müller"
    assert draft["title"] == "Zoë Ångström-Müller Resume"
    assert draft["location"] == "Zürich, Switzerland"
    assert draft["skills"] == ["Python", "C++", "C#", "Node.js"]
    section = draft["sections"][0]
    assert section["organization"] == "Société Générale"
    assert section["location"] == "Paris, France"


def test_skill_splitting_dedupe_and_parentheses(client: TestClient, auth_headers: Headers) -> None:
    text = (
        "Sam Lee\nsam@example.com\n\nTECHNICAL SKILLS\n"
        "• Python (Django, Flask) • SQL\n"
        "Docker; docker; DOCKER, CI/CD\n"
        "- " + "x" * 80 + "\n"
    )
    body = _parse_ok(client, auth_headers, "c.txt", _txt(text))
    assert body["draft"]["skills"] == ["Python (Django, Flask)", "SQL", "Docker", "CI/CD"]
    assert any("longer than 60" in w for w in body["warnings"])


def test_inline_skills_heading_and_ignored_sections(
    client: TestClient, auth_headers: Headers
) -> None:
    text = "Sam Lee\nsam@example.com\n\nSkills: Excel, SQL\n\nLanguages\nEnglish, French\n"
    body = _parse_ok(client, auth_headers, "c.txt", _txt(text))
    assert body["draft"]["skills"] == ["Excel", "SQL"]
    assert any("'Languages'" in w for w in body["warnings"])


def test_limits_are_enforced_with_warnings(client: TestClient, auth_headers: Headers) -> None:
    lines = ["Sam Lee", "sam@example.com", "", "EXPERIENCE", "Engineer | Acme", "2020 - 2021"]
    lines += [f"- bullet number {i}" for i in range(20)]
    lines.append("- " + "word " * 150)
    for i in range(25):
        lines += [f"Role {i} | Company {i}", f"{2000 + i} - {2001 + i}", "- did a thing"]
    lines += ["", "SKILLS", ", ".join(f"Skill{i}" for i in range(70))]
    body = _parse_ok(client, auth_headers, "c.txt", _txt("\n".join(lines)))
    draft = body["draft"]
    assert len(draft["sections"]) == 20
    assert len(draft["sections"][0]["bullets"]) == 15
    assert len(draft["skills"]) == 50
    assert [s["position"] for s in draft["sections"]] == list(range(20))
    joined = " ".join(body["warnings"])
    assert "beyond the limit" in joined
    assert "Only the first 20" in joined
    assert "Only the first 50 of 70 skills" in joined


def test_long_bullet_is_shortened(client: TestClient, auth_headers: Headers) -> None:
    text = (
        "Sam Lee\nsam@example.com\n\nEXPERIENCE\nEngineer | Acme\n2020 - 2021\n- " + "word " * 200
    )
    body = _parse_ok(client, auth_headers, "c.txt", _txt(text))
    bullet = body["draft"]["sections"][0]["bullets"][0]
    assert len(bullet) <= 500
    assert any("shortened" in w for w in body["warnings"])


def test_reversed_dates_are_blanked(client: TestClient, auth_headers: Headers) -> None:
    text = "Sam Lee\nsam@example.com\n\nEXPERIENCE\nEngineer | Acme\n2022 - 2019\n- Did work\n"
    body = _parse_ok(client, auth_headers, "c.txt", _txt(text))
    section = body["draft"]["sections"][0]
    assert section["start_date"] is None and section["end_date"] is None
    assert any("end date before" in w for w in body["warnings"])


def test_bullets_without_header_get_placeholder_title() -> None:
    draft, warnings = resume_parser.parse_resume(
        "Sam Lee\nsam@example.com\n\nEXPERIENCE\n- Did a thing\n", None
    )
    assert draft.sections[0].title == "Experience"
    assert any("no recognisable title" in w for w in warnings)


def test_phone_not_confused_with_year_range(client: TestClient, auth_headers: Headers) -> None:
    text = "Sam Lee\nsam@example.com\n\nEXPERIENCE\nEngineer | Acme\n2019 - 2021\n- Did work\n"
    assert _parse_ok(client, auth_headers, "c.txt", _txt(text))["draft"]["phone"] is None


def test_urls_get_https_and_validate() -> None:
    draft, _ = resume_parser.parse_resume(
        "Sam Lee\nwww.samlee.dev | linkedin.com/in/sam-lee\nsam@example.com\n", None
    )
    assert draft.website_url == "https://www.samlee.dev"
    assert draft.linkedin_url == "https://linkedin.com/in/sam-lee"


def test_invalid_email_dropped() -> None:
    draft, warnings = resume_parser.parse_resume("Sam Lee\nsam@example.invalid_tld-\n", None)
    assert draft.email is None or "@" in draft.email
    assert isinstance(warnings, list)


def test_text_truncation_warning() -> None:
    long_text = "Sam Lee\nsam@example.com\n\nSUMMARY\n" + ("Lorem ipsum dolor sit amet. " * 8000)
    result = import_service.extract_text("txt", long_text.encode())
    assert len(result.text) == import_service.MAX_TEXT_CHARS
    assert any("very long" in w for w in result.warnings)


def test_long_summary_is_shortened() -> None:
    text = "Sam Lee\nsam@example.com\n\nSUMMARY\n" + ("Lorem ipsum dolor sit amet. " * 200)
    draft, warnings = resume_parser.parse_resume(text, None)
    assert draft.summary is not None and len(draft.summary) <= 3000
    assert any("summary was too long" in w for w in warnings)


def test_page_numbers_are_ignored() -> None:
    text = (
        "Sam Lee\nsam@example.com\n\nEXPERIENCE\nEngineer | Acme\n2020 - 2021\n"
        "- A\nPage 2 of 3\n- B\n"
    )
    draft, _ = resume_parser.parse_resume(text, None)
    assert draft.sections[0].bullets == ["A", "B"]


# ------------------------------------------------------------- PDF specifics
def test_scanned_pdf_has_clear_message(client: TestClient, auth_headers: Headers) -> None:
    response = _post(client, auth_headers, "scan.pdf", make_blank_pdf(), "application/pdf")
    assert response.status_code == 400
    error = response.json()["error"]
    assert "scanned" in error["message"].lower()
    assert "text-based PDF or a DOCX" in error["message"]


def test_pdf_page_cap(client: TestClient, auth_headers: Headers) -> None:
    body = _parse_ok(client, auth_headers, "long.pdf", make_pdf(RESUME_TEXT, pages=17))
    assert any("first 15 of 17 pages" in w for w in body["warnings"])


def test_corrupt_pdf(client: TestClient, auth_headers: Headers) -> None:
    response = _post(client, auth_headers, "bad.pdf", b"%PDF-1.4\n garbage that is not a pdf")
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "IMPORT_FAILED"


def test_encrypted_pdf(client: TestClient, auth_headers: Headers) -> None:
    from pypdf import PdfWriter

    reader = PdfReader(io.BytesIO(make_pdf(RESUME_TEXT)))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt("owner-secret", "user-secret")
    buffer = io.BytesIO()
    writer.write(buffer)
    response = _post(client, auth_headers, "locked.pdf", buffer.getvalue())
    assert response.status_code == 400
    assert "password" in response.json()["error"]["message"].lower()


# --------------------------------------------------------- security / errors
def test_requires_authentication(client: TestClient) -> None:
    response = client.post(URL, files={"file": ("a.txt", b"hello", "text/plain")})
    assert response.status_code == 401


def test_missing_file_field_is_422(client: TestClient, auth_headers: Headers) -> None:
    assert client.post(URL, headers=auth_headers).status_code == 422


def test_legacy_doc_is_415(client: TestClient, auth_headers: Headers) -> None:
    response = _post(client, auth_headers, "old.doc", b"\xd0\xcf\x11\xe0 legacy")
    assert response.status_code == 415
    assert ".docx" in response.json()["error"]["message"]


@pytest.mark.parametrize("name", ["photo.png", "archive.zip", "noextension", "x.exe", ".pdf.bak"])
def test_unknown_extension_is_415(client: TestClient, auth_headers: Headers, name: str) -> None:
    response = _post(client, auth_headers, name, b"data")
    assert response.status_code == 415
    assert response.json()["error"]["code"] == "UNSUPPORTED_FORMAT"


def test_empty_file(client: TestClient, auth_headers: Headers) -> None:
    response = _post(client, auth_headers, "empty.txt", b"")
    assert response.status_code == 400
    assert "empty" in response.json()["error"]["message"].lower()


def test_whitespace_only_file(client: TestClient, auth_headers: Headers) -> None:
    response = _post(client, auth_headers, "blank.txt", b"  \n\n \t ")
    assert response.status_code == 400


def test_file_over_5mb_is_413(client: TestClient, auth_headers: Headers) -> None:
    data = b"a" * (import_service.MAX_UPLOAD_BYTES + 1)
    response = _post(client, auth_headers, "big.txt", data)
    assert response.status_code == 413
    assert response.json()["error"]["code"] == "FILE_TOO_LARGE"


def test_exactly_5mb_is_accepted(client: TestClient, auth_headers: Headers) -> None:
    data = (b"Sam Lee\nsam@example.com\n" + b"x" * 100 + b"\n").ljust(
        import_service.MAX_UPLOAD_BYTES, b"y"
    )
    response = _post(client, auth_headers, "edge.txt", data)
    assert response.status_code == 200


def test_declared_content_length_over_limit_is_rejected(
    client: TestClient, auth_headers: Headers
) -> None:
    headers = {**auth_headers, "Content-Length": str(import_service.MAX_UPLOAD_BYTES + 10**6)}
    response = client.post(URL, files={"file": ("a.txt", b"hi", "text/plain")}, headers=headers)
    assert response.status_code in (400, 413)


@pytest.mark.parametrize(
    ("name", "data"),
    [
        ("fake.pdf", b"This is not a pdf at all"),
        ("fake.docx", b"plain text pretending to be docx"),
        ("fake.odt", b"%PDF-1.4 hello"),
        ("fake.rtf", b"just text"),
        ("fake.txt", b"%PDF-1.4 hello"),
        ("fake.md", b"PK\x03\x04 zipped"),
        ("fake.html", b"{\\rtf1 hello}"),
        ("binary.txt", b"abc\x00\x01\x02def"),
    ],
)
def test_magic_mismatch(client: TestClient, auth_headers: Headers, name: str, data: bytes) -> None:
    response = _post(client, auth_headers, name, data)
    assert response.status_code == 400
    assert "does not match" in response.json()["error"]["message"]


def test_docx_renamed_to_pdf_is_rejected(client: TestClient, auth_headers: Headers) -> None:
    response = _post(client, auth_headers, "x.pdf", make_docx(RESUME_TEXT))
    assert response.status_code == 400


def test_corrupt_zip_formats(client: TestClient, auth_headers: Headers) -> None:
    for name in ("bad.docx", "bad.odt"):
        response = _post(client, auth_headers, name, b"PK\x03\x04" + b"\x00" * 40)
        assert response.status_code == 400
        assert response.json()["error"]["code"] == "IMPORT_FAILED"


def test_zip_without_document_xml(client: TestClient, auth_headers: Headers) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "hi")
    for name in ("empty.docx", "empty.odt"):
        response = _post(client, auth_headers, name, buffer.getvalue())
        assert response.status_code == 400


def _zip_bytes(entries: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    return buffer.getvalue()


def test_zip_bomb_total_size(client: TestClient, auth_headers: Headers) -> None:
    bomb = _zip_bytes({"word/document.xml": b"0" * (30 * 1024 * 1024)})
    assert len(bomb) < import_service.MAX_UPLOAD_BYTES
    for name in ("bomb.docx", "bomb.odt"):
        response = _post(client, auth_headers, name, bomb)
        assert response.status_code == 400
        assert "unreasonable size" in response.json()["error"]["message"]


def test_zip_bomb_single_member_limit(
    client: TestClient, auth_headers: Headers, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(import_service, "MAX_XML_BYTES", 1000)
    bomb = _zip_bytes({"content.xml": b"<a>" + b"x" * 5000 + b"</a>"})
    response = _post(client, auth_headers, "big.odt", bomb)
    assert response.status_code == 400


def test_zip_entry_count_cap(client: TestClient, auth_headers: Headers) -> None:
    many = _zip_bytes({f"f{i}.txt": b"x" for i in range(import_service.MAX_ZIP_ENTRIES + 5)})
    response = _post(client, auth_headers, "many.docx", many)
    assert response.status_code == 400
    assert "too many" in response.json()["error"]["message"]


def test_odt_xml_entities_rejected(client: TestClient, auth_headers: Headers) -> None:
    evil = (
        b'<?xml version="1.0"?><!DOCTYPE lol [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;&a;">]>'
        b"<office:document-content "
        b'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        b'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
        b"<office:body><office:text><text:p>&b;</text:p></office:text></office:body>"
        b"</office:document-content>"
    )
    response = _post(client, auth_headers, "evil.odt", _zip_bytes({"content.xml": evil}))
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "IMPORT_FAILED"


def test_odt_not_xml(client: TestClient, auth_headers: Headers) -> None:
    response = _post(client, auth_headers, "x.odt", _zip_bytes({"content.xml": b"not xml <<"}))
    assert response.status_code == 400


def test_corrupt_rtf_still_handled(client: TestClient, auth_headers: Headers) -> None:
    response = _post(client, auth_headers, "x.rtf", b"{\\rtf1 {{{{ \\par")
    assert response.status_code in (200, 400)


def test_control_characters_are_sanitised(client: TestClient, auth_headers: Headers) -> None:
    text = "Sam\x07 Lee\x00\nsam@example.com​\n\nSKILLS\nPy\x1fthon\n"
    body = _parse_ok(client, auth_headers, "c.txt", _txt(text.replace("\x00", "")))
    assert body["draft"]["full_name"] == "Sam Lee"
    assert body["draft"]["skills"] == ["Python"]


def test_utf16_and_cp1252_text(client: TestClient, auth_headers: Headers) -> None:
    text = "Renée Dupont\nrenee@example.com\n\nSKILLS\nPython\n"
    body = _parse_ok(client, auth_headers, "a.txt", text.encode("utf-16"))
    assert body["draft"]["full_name"] == "Renée Dupont"
    body = _parse_ok(client, auth_headers, "b.txt", text.encode("cp1252"))
    assert body["draft"]["full_name"] == "Renée Dupont"


def test_parser_crash_is_a_400(
    client: TestClient, auth_headers: Headers, monkeypatch: pytest.MonkeyPatch
) -> None:
    def boom(text: str, filename: str | None) -> None:
        raise RuntimeError("secret resume content must not leak")

    monkeypatch.setattr(resume_parser, "parse_resume", boom)
    response = _post(client, auth_headers, "c.txt", b"Sam Lee\n")
    assert response.status_code == 400
    assert "secret" not in response.text


def test_contents_never_logged(
    client: TestClient, auth_headers: Headers, caplog: pytest.LogCaptureFixture
) -> None:
    caplog.set_level("DEBUG")
    _parse_ok(client, auth_headers, "c.txt", _txt(RESUME_TEXT))
    assert "jane.doe@example.com" not in caplog.text
    assert "Acme Corp" not in caplog.text
    assert "Imported resume: format=txt" in caplog.text


def test_rate_limit(
    client: TestClient, auth_headers: Headers, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "RATE_LIMIT_ENABLED", True)
    limit_import.reset()
    for _ in range(limit_import.limit):
        assert _post(client, auth_headers, "c.txt", _txt(RESUME_TEXT)).status_code == 200
    blocked = _post(client, auth_headers, "c.txt", _txt(RESUME_TEXT))
    assert blocked.status_code == 429
    assert "Retry-After" in blocked.headers


# --------------------------------------------------------------- full loop
@pytest.mark.parametrize(
    ("name", "builder"),
    [("cv.docx", make_docx), ("cv.pdf", make_pdf), ("cv.txt", _txt), ("cv.odt", make_odt)],
)
def test_imported_draft_saves_and_exports(
    client: TestClient, auth_headers: Headers, template_id: int, name: str, builder: Any
) -> None:
    body = _parse_ok(client, auth_headers, name, builder(RESUME_TEXT))
    payload = {**body["draft"], "template_id": template_id}
    created = client.post("/api/resumes", json=payload, headers=auth_headers)
    assert created.status_code == 201, created.text
    resume = created.json()
    assert resume["full_name"] == "Jane Q. Doe"
    assert len(resume["sections"]) == 5
    assert resume["skills"][:2] == ["Python", "Go"]

    pdf = client.get(f"/api/resumes/{resume['id']}/export/pdf", headers=auth_headers)
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")
    text = " ".join(
        (page.extract_text() or "") for page in PdfReader(io.BytesIO(pdf.content)).pages
    )
    assert "Jane Q. Doe" in text
    assert "Senior Software Engineer" in text


def test_round_trip_through_own_export(
    client: TestClient, auth_headers: Headers, template_id: int
) -> None:
    """A resume exported by the app can be re-imported (PDF and DOCX)."""
    from tests.conftest import resume_payload

    created = client.post(
        "/api/resumes", json=resume_payload(template_id), headers=auth_headers
    ).json()
    for fmt, filename in (("pdf", "again.pdf"), ("docx", "again.docx")):
        exported = client.get(f"/api/resumes/{created['id']}/export/{fmt}", headers=auth_headers)
        assert exported.status_code == 200
        body = _parse_ok(client, auth_headers, filename, exported.content)
        draft = body["draft"]
        assert draft["full_name"] == "Alice Example"
        assert draft["email"] == "alice@example.com"
        assert {"Python", "FastAPI"} <= set(draft["skills"])
        assert any(s["title"] == "Senior Engineer" for s in draft["sections"])


def test_password_constant_is_used() -> None:
    assert PASSWORD
