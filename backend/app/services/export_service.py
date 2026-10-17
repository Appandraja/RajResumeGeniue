"""Resume export (PDF / DOCX / ODT / RTF / TXT / Markdown / HTML).

Pipeline: ``build_export_document`` turns a Resume + Template into a plain,
ORM-free ``ExportDocument``; format-specific renderers (``render_pdf``, ``render_docx``,
``render_odt``, ``render_rtf``, ``render_txt``, ``render_md``, ``render_html``) consume
only that model. New formats need a renderer plus an ``EXPORT_FORMATS`` entry.

ATS rules enforced by every renderer: single column, real text only (no images,
tables, text boxes, columns, headers/footers), standard section headings,
simple bullet characters, contact info as plain text at the top, entries in
reverse-chronological order (``end_date`` NULL = "Present"), and the accent
colour used only on section headings and their rules.

Font mapping (``Template.font_family`` -> built-in PDF font family, no embedding):
    Times New Roman, Georgia, Garamond, Cambria, any "serif" -> Times-Roman family
    Courier, Consolas, "mono"                                  -> Courier family
    Arial, Calibri, Helvetica, anything else                   -> Helvetica family
DOCX uses the template's font name directly (Word falls back sensibly).

Resume content is never logged.
"""

import io
import logging
import re
import textwrap
import unicodedata
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from datetime import date
from html import escape as html_escape
from typing import Any
from xml.sax.saxutils import escape

from reportlab.lib.colors import Color, HexColor, black
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)
from reportlab.platypus.flowables import Flowable

from app.exceptions import BadRequestError
from app.models.resume import Resume
from app.models.resume_section import SectionType
from app.models.template import Template

logger = logging.getLogger(__name__)

# ASCII hyphen: the WinAnsi bullet glyph extracts as a control char in some parsers.
_BULLET = "-"
_DEFAULT_ACCENT = "#000000"
_HEX_RE = re.compile(r"^#[0-9a-fA-F]{6}$")
_CTRL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")

# (regular, bold, italic) built-in PDF fonts.
_HELVETICA = ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique")
_TIMES = ("Times-Roman", "Times-Bold", "Times-Italic")
_COURIER = ("Courier", "Courier-Bold", "Courier-Oblique")

# Section order and headings; entry sections are keyed by SectionType.
_ENTRY_SECTIONS: tuple[tuple[SectionType, str], ...] = (
    (SectionType.experience, "Experience"),
    (SectionType.education, "Education"),
)
_TRAILING_ENTRY_SECTIONS: tuple[tuple[SectionType, str], ...] = (
    (SectionType.project, "Projects"),
    (SectionType.certification, "Certifications"),
)


# --------------------------------------------------------------------------- #
# Format-neutral document model
# --------------------------------------------------------------------------- #
@dataclass
class ExportEntry:
    title: str
    subtitle: str = ""  # "Organization, Location"
    dates: str = ""
    bullets: list[str] = field(default_factory=list)


@dataclass
class ExportSection:
    heading: str
    paragraph: str = ""  # used by Summary / Skills
    entries: list[ExportEntry] = field(default_factory=list)


@dataclass
class ExportDocument:
    title: str
    name: str
    contact: list[str]
    sections: list[ExportSection]
    font_family: str
    accent_color: str


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def _clean(value: object) -> str:
    """Return str with control characters removed and whitespace collapsed. None -> ''."""
    if value is None:
        return ""
    text = _CTRL_RE.sub("", str(value))
    return " ".join(text.split()) if "\n" in text or "\r" in text or "\t" in text else text.strip()


def _unrenderable_letters(texts: Iterable[str]) -> list[str]:
    """Letters/digits the built-in PDF fonts cannot show and that cannot be folded to ASCII.

    ``_pdf_safe`` would silently drop these (e.g. Cyrillic, CJK, "Ł"); callers should
    refuse the export instead so users never receive a resume with missing text.
    Symbols and emoji are not reported: dropping them loses no information.
    """
    found: dict[str, None] = {}
    for text in texts:
        for ch in text:
            try:
                ch.encode("cp1252")
                continue
            except UnicodeEncodeError:
                pass
            folded = unicodedata.normalize("NFKD", ch).encode("ascii", "ignore")
            if not folded and unicodedata.category(ch)[0] in "LMN":
                found[ch] = None
    return list(found)


def _document_texts(doc: "ExportDocument") -> Iterable[str]:
    yield doc.title
    yield doc.name
    yield from doc.contact
    for section in doc.sections:
        yield section.heading
        yield section.paragraph
        for entry in section.entries:
            yield entry.title
            yield entry.subtitle
            yield entry.dates
            yield from entry.bullets


def _pdf_safe(text: str) -> str:
    """Make text encodable by the built-in WinAnsi PDF fonts (fold or drop the rest)."""
    out: list[str] = []
    for ch in text:
        try:
            ch.encode("cp1252")
            out.append(ch)
            continue
        except UnicodeEncodeError:
            pass
        folded = unicodedata.normalize("NFKD", ch).encode("ascii", "ignore").decode("ascii")
        out.append(folded)
    return "".join(out)


def _fmt_date(value: date | None) -> str:
    return value.strftime("%b %Y") if value else ""


def _date_range(kind: SectionType, start: date | None, end: date | None) -> str:
    if start and end:
        return f"{_fmt_date(start)} - {_fmt_date(end)}"
    if start:
        # A certification with a single date is simply "issued"; others are ongoing.
        if kind == SectionType.certification:
            return _fmt_date(start)
        return f"{_fmt_date(start)} - Present"
    if end:
        return _fmt_date(end)
    return ""


def _sort_key(
    kind: SectionType, start: date | None, end: date | None, pos: int
) -> tuple[int, date, date, int]:
    has_dates = 1 if (start or end) else 0
    if end:
        eff_end = end
    elif start and kind != SectionType.certification:
        eff_end = date.max  # current
    else:
        eff_end = start or date.min
    return (has_dates, eff_end, start or date.min, -pos)


def _accent(value: str) -> str:
    return value if _HEX_RE.match(value or "") else _DEFAULT_ACCENT


def _font_family_key(font_family: str) -> str:
    name = (font_family or "").lower()
    serif_names = ("times", "georgia", "garamond", "cambria", "serif")
    if "sans" not in name and any(k in name for k in serif_names):
        return "times"
    if any(k in name for k in ("courier", "consolas", "mono")):
        return "courier"
    return "helvetica"


def pdf_fonts_for(font_family: str) -> tuple[str, str, str]:
    """Map a template font name to a (regular, bold, italic) built-in PDF font triple."""
    return {"times": _TIMES, "courier": _COURIER}.get(_font_family_key(font_family), _HELVETICA)


def safe_filename(title: str, ext: str) -> str:
    """ASCII-only, header-safe download filename derived from the resume title."""
    ascii_title = unicodedata.normalize("NFKD", title or "").encode("ascii", "ignore").decode()
    slug = re.sub(r"[^A-Za-z0-9]+", "_", ascii_title).strip("_")[:80] or "resume"
    return f"{slug}.{ext}"


def _dedupe(items: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for item in items:
        key = item.lower()
        if item and key not in seen:
            seen.add(key)
            out.append(item)
    return out


# --------------------------------------------------------------------------- #
# Document builder
# --------------------------------------------------------------------------- #
def build_export_document(resume: Resume, template: Template) -> ExportDocument:
    """Convert ORM objects to the format-neutral, ATS-ordered export model."""
    r: Any = resume
    contact = [
        c
        for c in (
            _clean(r.email),
            _clean(r.phone),
            _clean(r.location),
            _clean(r.linkedin_url),
            _clean(r.website_url),
        )
        if c
    ]
    sections: list[ExportSection] = []

    summary = _clean(r.summary)
    if summary:
        sections.append(ExportSection("Summary", paragraph=summary))

    def entries_for(kind: SectionType) -> list[ExportEntry]:
        rows = [s for s in (r.sections or []) if s.type == kind]
        rows.sort(
            key=lambda s: _sort_key(kind, s.start_date, s.end_date, s.position or 0),
            reverse=True,
        )
        result: list[ExportEntry] = []
        for s in rows:
            title = _clean(s.title)
            if not title:
                continue
            subtitle = ", ".join(p for p in (_clean(s.organization), _clean(s.location)) if p)
            bullets = [b for b in (_clean(b) for b in (s.bullets or [])) if b]
            result.append(
                ExportEntry(
                    title=title,
                    subtitle=subtitle,
                    dates=_date_range(kind, s.start_date, s.end_date),
                    bullets=bullets,
                )
            )
        return result

    def add_entries(kind: SectionType, heading: str) -> None:
        entries = entries_for(kind)
        if entries:
            sections.append(ExportSection(heading, entries=entries))

    for kind, heading in _ENTRY_SECTIONS:
        add_entries(kind, heading)

    skills = _dedupe(_clean(s) for s in (r.skills or []))
    if skills:
        sections.append(ExportSection("Skills", paragraph=", ".join(skills)))

    for kind, heading in _TRAILING_ENTRY_SECTIONS:
        add_entries(kind, heading)

    return ExportDocument(
        title=_clean(r.title) or "Resume",
        name=_clean(r.full_name),
        contact=contact,
        sections=sections,
        font_family=_clean(template.font_family),
        accent_color=_accent(_clean(template.accent_color)),
    )


# --------------------------------------------------------------------------- #
# PDF renderer (reportlab platypus)
# --------------------------------------------------------------------------- #
def _p(text: str) -> str:
    """Escape user text for reportlab Paragraph markup."""
    return escape(_pdf_safe(text))


def render_pdf(doc: ExportDocument) -> bytes:
    """Render an ExportDocument to a single-column, text-only PDF.

    Raises:
        BadRequestError: the resume contains letters the PDF fonts cannot display
            (e.g. Cyrillic or CJK). All other formats support them.
    """
    missing = _unrenderable_letters(_document_texts(doc))
    if missing:
        shown = " ".join(missing[:10])
        raise BadRequestError(
            f"PDF export cannot display these characters yet: {shown}. "
            "Please download DOCX, ODT, RTF, TXT, HTML or Markdown instead."
        )
    regular, bold, italic = pdf_fonts_for(doc.font_family)
    try:
        accent: Color = HexColor(doc.accent_color)
    except ValueError:
        accent = black

    body = ParagraphStyle("body", fontName=regular, fontSize=10, leading=13, alignment=TA_LEFT)
    name_style = ParagraphStyle(
        "name", parent=body, fontName=bold, fontSize=18, leading=22, spaceAfter=2
    )
    contact_style = ParagraphStyle("contact", parent=body, fontSize=9.5, leading=12, spaceAfter=4)
    heading_style = ParagraphStyle(
        "heading",
        parent=body,
        fontName=bold,
        fontSize=12,
        leading=15,
        textColor=accent,
        spaceBefore=10,
        spaceAfter=1,
        keepWithNext=1,
    )
    entry_title = ParagraphStyle(
        "entry_title", parent=body, fontName=bold, spaceBefore=5, keepWithNext=1
    )
    entry_meta = ParagraphStyle("entry_meta", parent=body, fontName=italic, keepWithNext=1)
    bullet_style = ParagraphStyle(
        "bullet", parent=body, leftIndent=14, bulletIndent=3, spaceBefore=1
    )

    story: list[Flowable] = []
    if doc.name:
        story.append(Paragraph(_p(doc.name), name_style))
    if doc.contact:
        story.append(Paragraph(_p(" | ".join(doc.contact)), contact_style))

    for section in doc.sections:
        story.append(Paragraph(_p(section.heading), heading_style))
        story.append(HRFlowable(width="100%", thickness=0.75, color=accent, spaceAfter=3))
        if section.paragraph:
            story.append(Paragraph(_p(section.paragraph), body))
        for entry in section.entries:
            head: list[Flowable] = [Paragraph(_p(entry.title), entry_title)]
            meta = " | ".join(m for m in (entry.subtitle, entry.dates) if m)
            has_bullets = bool(entry.bullets)
            if meta:
                head.append(Paragraph(_p(meta), entry_meta))
            if has_bullets:
                first = Paragraph(_p(entry.bullets[0]), bullet_style, bulletText=_BULLET)
                story.append(KeepTogether([*head, first]))
                for bullet in entry.bullets[1:]:
                    story.append(Paragraph(_p(bullet), bullet_style, bulletText=_BULLET))
            else:
                # Nothing follows; do not carry keepWithNext onto the next entry.
                for fl in head:
                    fl.keepWithNext = False
                story.extend(head)
        story.append(Spacer(1, 2))

    buffer = io.BytesIO()
    template = SimpleDocTemplate(
        buffer,
        pagesize=LETTER,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.7 * inch,
        title=_pdf_safe(doc.title),
        author=_pdf_safe(doc.name),
    )
    if not story:
        story.append(Paragraph(_p(doc.title), body))
    template.build(story)
    return buffer.getvalue()


def render_resume_pdf(resume: Resume, template: Template) -> bytes:
    """Render a resume as an ATS-friendly PDF and return its bytes."""
    return render_pdf(build_export_document(resume, template))


# --------------------------------------------------------------------------- #
# DOCX renderer (python-docx)
# --------------------------------------------------------------------------- #
def render_docx(doc: ExportDocument) -> bytes:
    """Render an ExportDocument to a single-column DOCX (no tables, no headers/footers)."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    word = Document()
    section = word.sections[0]
    section.left_margin = section.right_margin = Inches(0.75)
    section.top_margin = section.bottom_margin = Inches(0.7)

    font_name = doc.font_family or "Arial"
    accent = RGBColor.from_string(doc.accent_color.lstrip("#").upper())

    def pin_font(style_name: str) -> None:
        style = word.styles[style_name]
        style.font.name = font_name
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is not None:
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                key = qn(attr)
                if key in fonts.attrib:
                    del fonts.attrib[key]
            fonts.set(qn("w:eastAsia"), font_name)

    for style_name in ("Normal", "List Bullet", "Heading 1", "Title"):
        pin_font(style_name)
    word.styles["Normal"].font.size = Pt(10)
    word.styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)

    heading_style = word.styles["Heading 1"]
    heading_style.font.size = Pt(12)
    heading_style.font.bold = True
    heading_style.font.color.rgb = accent
    heading_style.paragraph_format.space_before = Pt(10)
    heading_style.paragraph_format.space_after = Pt(2)
    heading_style.paragraph_format.keep_with_next = True

    if doc.name:
        para = word.add_paragraph()
        run = para.add_run(doc.name)
        run.bold = True
        run.font.size = Pt(18)
    if doc.contact:
        word.add_paragraph(" | ".join(doc.contact))
    if not (doc.name or doc.contact or doc.sections):
        word.add_paragraph(doc.title)

    for sec in doc.sections:
        word.add_paragraph(sec.heading, style="Heading 1")
        if sec.paragraph:
            word.add_paragraph(sec.paragraph)
        for entry in sec.entries:
            head = word.add_paragraph()
            head.paragraph_format.space_before = Pt(4)
            head.paragraph_format.space_after = Pt(0)
            head.paragraph_format.keep_with_next = True
            head.add_run(entry.title).bold = True
            meta = " | ".join(m for m in (entry.subtitle, entry.dates) if m)
            if meta:
                head.add_run().add_break()
                head.add_run(meta).italic = True
            for bullet in entry.bullets:
                word.add_paragraph(bullet, style="List Bullet")

    core = word.core_properties
    core.title = doc.title
    core.author = doc.name

    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def render_resume_docx(resume: Resume, template: Template) -> bytes:
    """Render a resume as an ATS-friendly DOCX and return its bytes."""
    return render_docx(build_export_document(resume, template))


# --------------------------------------------------------------------------- #
# Shared helpers for the text-based renderers
# --------------------------------------------------------------------------- #
_FONT_NAME_RE = re.compile(r"[^A-Za-z0-9 \-]")


def _font_name(font_family: str) -> str:
    """Font name restricted to a safe charset (embedded in ODT/RTF/CSS)."""
    return " ".join(_FONT_NAME_RE.sub("", font_family or "").split()) or "Arial"


def _meta_line(entry: ExportEntry) -> str:
    return " | ".join(m for m in (entry.subtitle, entry.dates) if m)


def _is_empty(doc: ExportDocument) -> bool:
    return not (doc.name or doc.contact or doc.sections)


# --------------------------------------------------------------------------- #
# ODT renderer (odfpy)
# --------------------------------------------------------------------------- #
def render_odt(doc: ExportDocument) -> bytes:
    """Render an ExportDocument to a single-column ODT (paragraphs, headings, list items)."""
    from odf import dc
    from odf.opendocument import OpenDocumentText
    from odf.style import (
        FontFace,
        ListLevelProperties,
        ParagraphProperties,
        Style,
        TextProperties,
    )
    from odf.text import H, List, ListItem, ListLevelStyleBullet, ListStyle
    from odf.text import P as OdfP

    font = _font_name(doc.font_family)
    odt = OpenDocumentText()
    odt.fontfacedecls.addElement(FontFace(name=font, fontfamily=f"'{font}'"))

    def para_style(name: str, text_props: dict[str, str], para_props: dict[str, str]) -> Style:
        style = Style(name=name, family="paragraph")
        style.addElement(TextProperties(fontname=font, **text_props))
        if para_props:
            style.addElement(ParagraphProperties(**para_props))
        odt.styles.addElement(style)
        return style

    body = para_style("ResumeBody", {"fontsize": "10pt"}, {})
    name_style = para_style(
        "ResumeName", {"fontsize": "18pt", "fontweight": "bold"}, {"marginbottom": "0.03in"}
    )
    heading_style = para_style(
        "ResumeHeading",
        {"fontsize": "12pt", "fontweight": "bold", "color": doc.accent_color},
        {"margintop": "0.14in", "marginbottom": "0.03in", "keepwithnext": "always"},
    )
    entry_style = para_style(
        "ResumeEntryTitle",
        {"fontsize": "10pt", "fontweight": "bold"},
        {"margintop": "0.06in", "keepwithnext": "always"},
    )
    meta_style = para_style("ResumeEntryMeta", {"fontsize": "10pt", "fontstyle": "italic"}, {})

    bullets = ListStyle(name="ResumeBullets")
    level = ListLevelStyleBullet(level=1, bulletchar="-", stylename="ResumeBody")
    level.addElement(ListLevelProperties(spacebefore="0.15in", minlabelwidth="0.15in"))
    bullets.addElement(level)
    odt.automaticstyles.addElement(bullets)

    def add_para(text: str, style: Style) -> None:
        odt.text.addElement(OdfP(stylename=style, text=text))

    if doc.name:
        add_para(doc.name, name_style)
    if doc.contact:
        add_para(" | ".join(doc.contact), body)
    if _is_empty(doc):
        add_para(doc.title, body)

    for sec in doc.sections:
        odt.text.addElement(H(outlinelevel=1, stylename=heading_style, text=sec.heading))
        if sec.paragraph:
            add_para(sec.paragraph, body)
        for entry in sec.entries:
            add_para(entry.title, entry_style)
            meta = _meta_line(entry)
            if meta:
                add_para(meta, meta_style)
            if entry.bullets:
                items = List(stylename=bullets)
                for bullet in entry.bullets:
                    item = ListItem()
                    item.addElement(OdfP(stylename=body, text=bullet))
                    items.addElement(item)
                odt.text.addElement(items)

    odt.meta.addElement(dc.Title(text=doc.title))
    if doc.name:
        odt.meta.addElement(dc.Creator(text=doc.name))

    buffer = io.BytesIO()
    odt.write(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# RTF renderer (hand-written)
# --------------------------------------------------------------------------- #
def _rtf(text: str) -> str:
    r"""Escape text for RTF: ``\``, ``{``, ``}`` and every non-ASCII char as ``\uN?``.

    Code points above U+FFFF are emitted as two UTF-16 units (signed 16-bit values).
    """
    out: list[str] = []
    for ch in text:
        code = ord(ch)
        if ch in "\\{}":
            out.append("\\" + ch)
        elif 32 <= code < 127:
            out.append(ch)
        elif code >= 128:
            units = [code]
            if code > 0xFFFF:
                code -= 0x10000
                units = [0xD800 + (code >> 10), 0xDC00 + (code & 0x3FF)]
            for unit in units:
                out.append(f"\\u{unit - 65536 if unit > 32767 else unit}?")
        # remaining ASCII controls are already stripped by ``_clean``
    return "".join(out)


def render_rtf(doc: ExportDocument) -> bytes:
    """Render an ExportDocument to a single-column RTF (plain paragraphs, no tables)."""
    font = _font_name(doc.font_family)
    red, green, blue = (int(doc.accent_color[i : i + 2], 16) for i in (1, 3, 5))
    parts: list[str] = [
        "{\\rtf1\\ansi\\ansicpg1252\\uc1\\deff0",
        f"{{\\fonttbl{{\\f0\\fnil\\fcharset0 {_rtf(font)};}}}}",
        f"{{\\colortbl;\\red{red}\\green{green}\\blue{blue};}}",
        f"{{\\info{{\\title {_rtf(doc.title)}}}{{\\author {_rtf(doc.name)}}}}}",
        "\\viewkind1\\paperw12240\\paperh15840\\margl1080\\margr1080\\margt1008\\margb1008",
        "\\f0\\fs20",
    ]

    if doc.name:
        parts.append(f"\\pard\\sa40\\b\\fs36 {_rtf(doc.name)}\\b0\\fs20\\par")
    if doc.contact:
        parts.append(f"\\pard\\sa60 {_rtf(' | '.join(doc.contact))}\\par")
    if _is_empty(doc):
        parts.append(f"\\pard\\sa60 {_rtf(doc.title)}\\par")

    for sec in doc.sections:
        parts.append(
            f"\\pard\\sb200\\sa60\\keepn\\cf1\\b\\fs24 {_rtf(sec.heading)}\\b0\\cf0\\fs20\\par"
        )
        if sec.paragraph:
            parts.append(f"\\pard\\sa60 {_rtf(sec.paragraph)}\\par")
        for entry in sec.entries:
            parts.append(f"\\pard\\sb80\\sa0\\keepn\\b {_rtf(entry.title)}\\b0\\par")
            meta = _meta_line(entry)
            if meta:
                parts.append(f"\\pard\\sa0\\i {_rtf(meta)}\\i0\\par")
            for bullet in entry.bullets:
                parts.append(f"\\pard\\sa20\\li360\\fi-180 - {_rtf(bullet)}\\par")
    parts.append("}")
    return "\r\n".join(parts).encode("ascii")


# --------------------------------------------------------------------------- #
# Plain text renderer
# --------------------------------------------------------------------------- #
_TXT_WIDTH = 80


def render_txt(doc: ExportDocument) -> bytes:
    """Render an ExportDocument to ATS-ideal UTF-8 plain text."""

    def wrap(text: str, first: str = "", rest: str = "") -> list[str]:
        return textwrap.wrap(
            text,
            width=_TXT_WIDTH,
            initial_indent=first,
            subsequent_indent=rest,
            break_long_words=False,
            break_on_hyphens=False,
        )

    lines: list[str] = []
    if doc.name:
        lines.extend(wrap(doc.name))
    if doc.contact:
        lines.extend(wrap(" | ".join(doc.contact)))
    if _is_empty(doc):
        lines.extend(wrap(doc.title))

    for sec in doc.sections:
        if lines:
            lines.append("")
        lines.append(sec.heading.upper())
        if sec.paragraph:
            lines.extend(wrap(sec.paragraph))
        for index, entry in enumerate(sec.entries):
            if index or sec.paragraph:
                lines.append("")
            lines.extend(wrap(entry.title))
            meta = _meta_line(entry)
            if meta:
                lines.extend(wrap(meta))
            for bullet in entry.bullets:
                lines.extend(wrap(bullet, "- ", "  "))
    return ("\n".join(lines) + "\n").encode("utf-8")


# --------------------------------------------------------------------------- #
# Markdown renderer
# --------------------------------------------------------------------------- #
_MD_SPECIAL_RE = re.compile(r"([\\`*_\[\]<>#|~&])")
_MD_LEADING_RE = re.compile(r"^(?:([+\-=])|(\d+)([.)]))")


def _md(text: str) -> str:
    """Escape Markdown-significant characters in user text."""
    escaped = _MD_SPECIAL_RE.sub(r"\\\1", text)
    match = _MD_LEADING_RE.match(escaped)
    if match is None:
        return escaped
    if match.group(1):  # list / setext marker at the start of a block
        return "\\" + escaped
    return escaped[: match.end(2)] + "\\" + escaped[match.end(2) :]  # "1." -> "1\."


def render_md(doc: ExportDocument) -> bytes:
    """Render an ExportDocument to Markdown (# name, ## sections, '- ' bullets)."""
    blocks: list[str] = []
    if doc.name:
        blocks.append(f"# {_md(doc.name)}")
    if doc.contact:
        blocks.append(_md(" | ".join(doc.contact)))
    if _is_empty(doc):
        blocks.append(f"# {_md(doc.title)}")

    for sec in doc.sections:
        blocks.append(f"## {_md(sec.heading)}")
        if sec.paragraph:
            blocks.append(_md(sec.paragraph))
        for entry in sec.entries:
            head = f"### {_md(entry.title)}"
            meta = _meta_line(entry)
            if meta:
                head += f"\n\n*{_md(meta)}*"
            blocks.append(head)
            if entry.bullets:
                blocks.append("\n".join(f"- {_md(b)}" for b in entry.bullets))
    return ("\n\n".join(blocks) + "\n").encode("utf-8")


# --------------------------------------------------------------------------- #
# HTML renderer
# --------------------------------------------------------------------------- #
_CSS_GENERIC = {"times": "serif", "courier": "monospace", "helvetica": "sans-serif"}


def render_html(doc: ExportDocument) -> bytes:
    """Render an ExportDocument to a self-contained, script-free HTML5 document."""

    def h(text: str) -> str:
        return html_escape(text, quote=True)

    font = _font_name(doc.font_family)
    generic = _CSS_GENERIC[_font_family_key(doc.font_family)]
    accent = doc.accent_color
    out: list[str] = [
        "<!DOCTYPE html>",
        '<html lang="en">',
        "<head>",
        '<meta charset="utf-8">',
        '<meta name="viewport" content="width=device-width, initial-scale=1">',
        f"<title>{h(doc.title)}</title>",
        "<style>",
        f'body {{ font-family: "{font}", {generic}; font-size: 11pt; line-height: 1.4;'
        " color: #000; background: #fff; max-width: 48rem; margin: 2rem auto; padding: 0 1rem; }",
        "h1 { font-size: 1.7em; margin: 0 0 0.2em; }",
        f"h2 {{ font-size: 1.15em; color: {accent};"
        f" border-bottom: 1px solid {accent}; margin: 1.2em 0 0.4em; }}",
        "h3 { font-size: 1em; margin: 0.7em 0 0; }",
        "p { margin: 0.2em 0; }",
        "p.meta { font-style: italic; }",
        "ul { margin: 0.2em 0; padding-left: 1.4em; }",
        "</style>",
        "</head>",
        "<body>",
        f"<h1>{h(doc.name or doc.title)}</h1>",
    ]
    if doc.contact:
        out.append(f'<p class="contact">{h(" | ".join(doc.contact))}</p>')
    for sec in doc.sections:
        out.append("<section>")
        out.append(f"<h2>{h(sec.heading)}</h2>")
        if sec.paragraph:
            out.append(f"<p>{h(sec.paragraph)}</p>")
        for entry in sec.entries:
            out.append(f"<h3>{h(entry.title)}</h3>")
            meta = _meta_line(entry)
            if meta:
                out.append(f'<p class="meta">{h(meta)}</p>')
            if entry.bullets:
                out.append("<ul>")
                out.extend(f"<li>{h(b)}</li>" for b in entry.bullets)
                out.append("</ul>")
        out.append("</section>")
    out.extend(["</body>", "</html>", ""])
    return "\n".join(out).encode("utf-8")


# --------------------------------------------------------------------------- #
# Format registry
# --------------------------------------------------------------------------- #
@dataclass(frozen=True)
class ExportFormat:
    """One downloadable format: ``key`` doubles as the URL segment and file extension."""

    key: str
    media_type: str
    render: Callable[[ExportDocument], bytes]


EXPORT_FORMATS: dict[str, ExportFormat] = {
    f.key: f
    for f in (
        ExportFormat("pdf", "application/pdf", render_pdf),
        ExportFormat(
            "docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            render_docx,
        ),
        ExportFormat("odt", "application/vnd.oasis.opendocument.text", render_odt),
        ExportFormat("rtf", "application/rtf", render_rtf),
        ExportFormat("txt", "text/plain; charset=utf-8", render_txt),
        ExportFormat("md", "text/markdown; charset=utf-8", render_md),
        ExportFormat("html", "text/html; charset=utf-8", render_html),
    )
}
