"""Resume import: safe in-memory text extraction for uploaded files.

Everything is processed in memory (nothing is written to disk). File contents and
extracted text are never logged; only the format, sizes and counts are.
"""

import io
import logging
import posixpath
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from typing import Final
from xml.etree.ElementTree import Element

from defusedxml import ElementTree as SafeET
from defusedxml.common import DefusedXmlException

from app.exceptions import AppException, BadRequestError
from app.schemas.importer import ImportResult, ImportSourceFormat
from app.services import resume_parser

logger = logging.getLogger(__name__)

MAX_UPLOAD_BYTES: Final = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES: Final = 25 * 1024 * 1024
MAX_ZIP_ENTRIES: Final = 2000
MAX_PDF_PAGES: Final = 15
MAX_TEXT_CHARS: Final = 200_000
MAX_XML_BYTES: Final = 20 * 1024 * 1024

_EXTENSIONS: Final[dict[str, ImportSourceFormat]] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".odt": "odt",
    ".rtf": "rtf",
    ".txt": "txt",
    ".md": "md",
    ".markdown": "md",
    ".html": "html",
    ".htm": "html",
}
_UNSUPPORTED_MESSAGE: Final = (
    "Unsupported file type. Upload a PDF, DOCX, ODT, RTF, TXT, Markdown or HTML resume."
)
_ZIP_MAGIC: Final = b"PK\x03\x04"
_PRIVATE_USE_RE: Final = re.compile("[-]")
_CONTROL_RE: Final = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f​-‏‪-‮﻿]")
_BLANK_RUN_RE: Final = re.compile(r"\n{3,}")
_ODT_TEXT_NS: Final = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
_ODT_OFFICE_NS: Final = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"


class UnsupportedFormatError(AppException):
    status_code = 415
    code = "UNSUPPORTED_FORMAT"


class FileTooLargeError(AppException):
    status_code = 413
    code = "FILE_TOO_LARGE"


class ImportParseError(BadRequestError):
    code = "IMPORT_FAILED"


@dataclass
class Extracted:
    text: str
    warnings: list[str] = field(default_factory=list)


# ----------------------------------------------------------------- detection
def detect_format(filename: str | None) -> ImportSourceFormat:
    """Map a filename to a supported format or raise 415."""
    name = (filename or "").strip().lower()
    ext = posixpath.splitext(name)[1]
    if ext == ".doc":
        raise UnsupportedFormatError(
            "Legacy .doc files are not supported. Please save your resume as .docx "
            "(or PDF) and upload it again."
        )
    fmt = _EXTENSIONS.get(ext)
    if fmt is None:
        raise UnsupportedFormatError(_UNSUPPORTED_MESSAGE)
    return fmt


def _mismatch() -> ImportParseError:
    return ImportParseError(
        "The file content does not match its extension. Check that it is a real resume "
        "file in the format you selected."
    )


def verify_magic(fmt: ImportSourceFormat, data: bytes) -> None:
    """Raise a 400 when the leading bytes contradict the claimed format."""
    head = data[:1024].lstrip()
    if fmt == "pdf":
        if not head.startswith(b"%PDF"):
            raise _mismatch()
    elif fmt in ("docx", "odt"):
        if not data.startswith(_ZIP_MAGIC):
            raise _mismatch()
    elif fmt == "rtf":
        if not head.startswith(b"{\\rtf"):
            raise _mismatch()
    else:
        is_utf16 = data.startswith((b"\xff\xfe", b"\xfe\xff"))
        if head.startswith((b"%PDF", _ZIP_MAGIC, b"{\\rtf")) or (
            b"\x00" in data[:4096] and not is_utf16
        ):
            raise _mismatch()


# --------------------------------------------------------------- sanitising
def sanitize_text(text: str) -> str:
    """Normalise unicode/newlines and strip control characters."""
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace(" ", "\n").replace(" ", "\n")
    text = text.replace("\x0b", "\n").replace("\x0c", "\n")
    text = text.replace("\t", "  ")
    text = _PRIVATE_USE_RE.sub("•", text)  # symbol-font bullets
    text = _CONTROL_RE.sub("", text)
    lines = [line.rstrip() for line in text.split("\n")]
    return _BLANK_RUN_RE.sub("\n\n", "\n".join(lines)).strip()


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


# --------------------------------------------------------------------- zip
def _open_zip(data: bytes) -> zipfile.ZipFile:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        infos = archive.infolist()
    except (zipfile.BadZipFile, ValueError, OSError) as exc:
        raise ImportParseError("The file is corrupt or not a valid document.") from exc
    if len(infos) > MAX_ZIP_ENTRIES:
        raise ImportParseError("The document contains too many internal files.")
    if sum(i.file_size for i in infos) > MAX_UNCOMPRESSED_BYTES:
        raise ImportParseError("The document expands to an unreasonable size and was rejected.")
    return archive


def _read_member(archive: zipfile.ZipFile, name: str, limit: int = MAX_XML_BYTES) -> bytes:
    try:
        with archive.open(name) as handle:
            content = handle.read(limit + 1)
    except KeyError as exc:
        raise ImportParseError("The document is missing expected content.") from exc
    except (zipfile.BadZipFile, RuntimeError, OSError, ValueError) as exc:
        raise ImportParseError("The file is corrupt or not a valid document.") from exc
    if len(content) > limit:
        raise ImportParseError("The document expands to an unreasonable size and was rejected.")
    return content


# --------------------------------------------------------------------- PDF
def _extract_pdf(data: bytes) -> Extracted:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ImportParseError("This PDF is password protected. Remove the password first.")
        total = len(reader.pages)
        parts: list[str] = []
        size = 0
        for index in range(min(total, MAX_PDF_PAGES)):
            page_text = reader.pages[index].extract_text() or ""
            parts.append(page_text)
            size += len(page_text)
            if size > MAX_TEXT_CHARS * 2:
                break
    except AppException:
        raise
    except Exception as exc:
        raise ImportParseError("Could not read this PDF. The file may be corrupt.") from exc
    warnings: list[str] = []
    if total > MAX_PDF_PAGES:
        warnings.append(f"Only the first {MAX_PDF_PAGES} of {total} pages were read.")
    text = "\n".join(parts)
    if not text.strip():
        raise ImportParseError(
            "This PDF has no selectable text. Scanned or image-only resumes are not "
            "supported yet; please upload a text-based PDF or a DOCX file."
        )
    return Extracted(text, warnings)


# -------------------------------------------------------------------- DOCX
def _extract_docx(data: bytes) -> Extracted:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _open_zip(data).close()
    lines: list[str] = []
    try:
        document = Document(io.BytesIO(data))

        def add_paragraph(paragraph: Paragraph) -> None:
            text = paragraph.text.strip()
            if not text:
                lines.append("")
                return
            style = paragraph.style.name if paragraph.style is not None else ""
            ppr = paragraph._p.pPr
            is_list = (style or "").lower().startswith("list") or (
                ppr is not None and ppr.numPr is not None
            )
            lines.append(f"- {text}" if is_list else text)

        def add_table(table: Table) -> None:
            seen: set[int] = set()
            for row in table.rows:
                for cell in row.cells:
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    for paragraph in cell.paragraphs:
                        add_paragraph(paragraph)
                lines.append("")

        for child in document.element.body.iterchildren():
            tag = str(child.tag).rsplit("}", 1)[-1]
            if tag == "p":
                add_paragraph(Paragraph(child, document))
            elif tag == "tbl":
                add_table(Table(child, document))
    except AppException:
        raise
    except Exception as exc:
        raise ImportParseError("Could not read this DOCX file. It may be corrupt.") from exc
    return Extracted("\n".join(lines))


# --------------------------------------------------------------------- ODT
def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _odt_inline(element: Element) -> str:
    parts: list[str] = [element.text or ""]
    for child in element:
        name = _local(child.tag)
        if name == "s":
            count = child.get(f"{_ODT_TEXT_NS}c", "1")
            parts.append(" " * (int(count) if count.isdigit() and int(count) < 50 else 1))
        elif name == "tab":
            parts.append("  ")
        elif name == "line-break":
            parts.append("\n")
        else:
            parts.append(_odt_inline(child))
        parts.append(child.tail or "")
    return "".join(parts)


def _odt_blocks(element: Element, in_list: bool, out: list[str]) -> None:
    name = _local(element.tag)
    if name in ("h", "p"):
        text = _odt_inline(element).strip()
        out.append(f"- {text}" if in_list and text else text)
        return
    child_in_list = in_list or name == "list-item"
    for child in element:
        _odt_blocks(child, child_in_list, out)
    if name in ("table-row", "table"):
        out.append("")


def _extract_odt(data: bytes) -> Extracted:
    archive = _open_zip(data)
    try:
        xml = _read_member(archive, "content.xml")
    finally:
        archive.close()
    try:
        root = SafeET.fromstring(xml, forbid_dtd=True)
    except (DefusedXmlException, SafeET.ParseError, ValueError) as exc:
        raise ImportParseError("Could not read this ODT file. It may be corrupt.") from exc
    out: list[str] = []
    for body in root.iter(f"{_ODT_OFFICE_NS}text"):
        _odt_blocks(body, False, out)
    return Extracted("\n".join(out))


# --------------------------------------------------------------------- RTF
def _extract_rtf(data: bytes) -> Extracted:
    from striprtf.striprtf import rtf_to_text

    try:
        text = rtf_to_text(data.decode("latin-1"), errors="ignore")
    except Exception as exc:
        raise ImportParseError("Could not read this RTF file. It may be corrupt.") from exc
    return Extracted(text)


# -------------------------------------------------------------------- HTML
_BLOCK_TAGS: Final = frozenset(
    "address article aside blockquote body br dd div dl dt fieldset footer form h1 h2 h3 h4 h5 h6 "
    "header hr li main nav ol p pre section table tbody td tfoot th thead tr ul".split()
)
_SKIP_TAGS: Final = frozenset({"script", "style", "head", "noscript", "template", "svg"})


class _TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in _SKIP_TAGS:
            self._skip += 1
        elif tag == "li":
            self.parts.append("\n- ")
        elif tag in _BLOCK_TAGS:
            self.parts.append("\n")

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in _BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in _SKIP_TAGS:
            self._skip = max(0, self._skip - 1)
        elif tag in _BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


def _extract_html(data: bytes) -> Extracted:
    parser = _TextExtractor()
    try:
        parser.feed(_decode(data))
        parser.close()
    except Exception as exc:
        raise ImportParseError("Could not read this HTML file.") from exc
    text = re.sub(r"[ \t]+", " ", "".join(parser.parts))
    return Extracted(re.sub(r"\n\s*\n", "\n", text))


# ---------------------------------------------------------------- Markdown
_MD_LINK_RE: Final = re.compile(r"!?\[([^\]]*)\]\(([^)\s]+)[^)]*\)")
_MD_EMPH_RE: Final = re.compile(r"(\*\*|__|`)")
_MD_HEADING_RE: Final = re.compile(r"^\s{0,3}#{1,6}\s*")
_MD_RULE_RE: Final = re.compile(r"^\s*([-*_=])(\s*\1){2,}\s*$")
_MD_BULLET_RE: Final = re.compile(r"^(\s*)[*+]\s+")


def _extract_markdown(data: bytes) -> Extracted:
    lines: list[str] = []
    for raw in _decode(data).splitlines():
        if _MD_RULE_RE.match(raw):
            lines.append("")
            continue
        line = _MD_HEADING_RE.sub("", raw)
        line = _MD_LINK_RE.sub(lambda m: f"{m.group(1)} {m.group(2)}".strip(), line)
        line = _MD_EMPH_RE.sub("", line)
        line = _MD_BULLET_RE.sub(r"\1- ", line)
        lines.append(line)
    return Extracted("\n".join(lines))


# ----------------------------------------------------------- orchestration
def extract_text(fmt: ImportSourceFormat, data: bytes) -> Extracted:
    """Extract sanitised plain text from `data` for the given (verified) format."""
    if fmt == "pdf":
        extracted = _extract_pdf(data)
    elif fmt == "docx":
        extracted = _extract_docx(data)
    elif fmt == "odt":
        extracted = _extract_odt(data)
    elif fmt == "rtf":
        extracted = _extract_rtf(data)
    elif fmt == "html":
        extracted = _extract_html(data)
    elif fmt == "md":
        extracted = _extract_markdown(data)
    else:
        extracted = Extracted(_decode(data))
    text = sanitize_text(extracted.text)
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS]
        extracted.warnings.append(
            f"The document is very long; only the first {MAX_TEXT_CHARS:,} characters were read."
        )
    if not text.strip():
        raise ImportParseError("No text could be found in this file. Is it empty?")
    return Extracted(text, extracted.warnings)


def import_resume(filename: str | None, data: bytes) -> ImportResult:
    """Validate, extract and heuristically parse an uploaded resume into a draft."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise FileTooLargeError("The file is too large. The maximum size is 5 MB.")
    fmt = detect_format(filename)
    if not data:
        raise ImportParseError("The uploaded file is empty.")
    verify_magic(fmt, data)
    extracted = extract_text(fmt, data)
    try:
        draft, warnings = resume_parser.parse_resume(extracted.text, filename)
    except Exception as exc:  # the parser must never surface a 500
        logger.exception("Resume parser failed (format=%s)", fmt)
        raise ImportParseError("We could not understand the structure of this resume.") from exc
    logger.info(
        "Imported resume: format=%s bytes=%d chars=%d sections=%d skills=%d warnings=%d",
        fmt,
        len(data),
        len(extracted.text),
        len(draft.sections),
        len(draft.skills),
        len(extracted.warnings) + len(warnings),
    )
    return ImportResult(
        draft=draft,
        warnings=[*extracted.warnings, *warnings],
        source_format=fmt,
        char_count=len(extracted.text),
    )
